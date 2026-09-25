"""Import the three AHR Excel workbooks into the inventory database.

Run once to seed, or re-run to upsert (matched by item_code). Handles the
quirks found while studying the source files:
  * real headers are on row 3, data starts row 4
  * Description packs several fields separated by '\\' and ':'
  * UOM casing is inconsistent (Pcs / PCS / pcs / Pcs.)
  * Machine field is free-text and coarse; the Machine list is the controlled
    vocabulary the dropdown will use going forward
Usage:
  python -m scripts.import_excel --dir /path/to/folder   (defaults to ./source)
"""
import argparse
import re
from datetime import datetime
from pathlib import Path

import openpyxl
from sqlalchemy import select

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from app.db import engine, init_db, SessionLocal, Session
from app.models import Machine, Item, StockTxn

UOM_MAP = {
    "pcs": "Pcs", "pc": "Pcs", "set": "Set", "uni": "Unit", "unit": "Unit",
    "pail": "Pail", "roll": "Roll", "kgs": "Kg", "kg": "Kg", "lites": "Litre",
    "litre": "Litre", "l": "Litre", "drum": "Drum", "meters": "Metre",
    "metre": "Metre", "m": "Metre", "box": "Box",
}


def clean_uom(v):
    if not v:
        return "Pcs"
    key = str(v).strip().strip(".").lower()
    return UOM_MAP.get(key, str(v).strip())


def parse_description(desc: str):
    """'GROUP \\ NAME : PARTNO \\ BRAND' -> (part_name, part_number, brand)."""
    if not desc:
        return "", "", ""
    parts = [p.strip() for p in str(desc).split("\\") if p.strip()]
    brand = ""
    if len(parts) >= 3:
        brand = parts[-1]
        core = parts[1]
    elif len(parts) == 2:
        core = parts[1]
    else:
        core = parts[0]
    name, pn = core, ""
    if ":" in core:
        name, pn = core.split(":", 1)
    return name.strip(), pn.strip(), brand.strip()


def cat_of(code: str) -> str:
    m = re.match(r"([A-Za-z]+)", str(code or ""))
    return m.group(1).upper() if m else ""


def as_num(v, default=0.0):
    try:
        if v in (None, "", "#DIV/0!"):
            return default
        return float(v)
    except (ValueError, TypeError):
        return default


def as_dt(v):
    return v if isinstance(v, datetime) else None


def upsert_item(s: Session, code, **fields):
    it = s.exec(select(Item).where(Item.item_code == code)).first()
    if it:
        for k, v in fields.items():
            setattr(it, k, v)
    else:
        it = Item(item_code=code, **fields)
        s.add(it)
    return it


def import_machines(path: Path, s: Session):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Machine list"]
    existing = {m.name for m in s.exec(select(Machine)).all()}
    n = 0
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if name and str(name).strip() and str(name).strip() not in existing:
            s.add(Machine(name=str(name).strip()))
            existing.add(str(name).strip())
            n += 1
    s.commit()
    print(f"  machines: +{n} (total {len(existing)})")


def enrich_from_itemcode(path: Path):
    """Build code -> {brand, model, history} from the 'Item code' master sheet."""
    out = {}
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Item code" not in wb.sheetnames:
        return out
    ws = wb["Item code"]
    # cols: D=4 code, F=6 history(th), G=7 model, H=8 brand
    for r in range(4, ws.max_row + 1):
        code = ws.cell(r, 4).value
        if not code:
            continue
        out[str(code).strip()] = {
            "history": str(ws.cell(r, 6).value or "").strip(),
            "model": str(ws.cell(r, 7).value or "").strip(),
            "brand": str(ws.cell(r, 8).value or "").strip(),
        }
    return out


def read_sheet(path: Path, sheet: str, enrich=None, kind="spare"):
    """Read one sheet into a list of row dicts (no database writes)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    enrich = enrich or {}
    if kind == "spare":
        C = dict(box=4, level=5, code=6, desc=7, machine=8, uom=9, stock=10,
                 out_year=49, mn=54, mx=53, lead=57, receipt=58, price=62)
    else:  # tool/facility
        C = dict(box=4, level=5, code=6, desc=7, machine=8, uom=9, stock=10,
                 out_year=49, mn=None, mx=None, lead=None, receipt=None, price=None)
    rows = []
    for r in range(4, ws.max_row + 1):
        code = ws.cell(r, C["code"]).value
        if not code or not str(code).strip():
            continue
        code = str(code).strip()
        desc = str(ws.cell(r, C["desc"]).value or "").strip()
        name, pn, brand = parse_description(desc)
        info = enrich.get(code, {})
        rows.append(dict(
            code=code, sheet=sheet, row=r, desc=desc,
            name=name or info.get("history", ""), pn=pn, brand=brand or info.get("brand", ""),
            machine=str(ws.cell(r, C["machine"]).value or "").strip(),
            uom=clean_uom(ws.cell(r, C["uom"]).value),
            box=str(ws.cell(r, C["box"]).value or "").strip(),
            level=str(ws.cell(r, C["level"]).value or "").strip(),
            stock=as_num(ws.cell(r, C["stock"]).value),
            out_year=as_num(ws.cell(r, C["out_year"]).value) if C["out_year"] else 0.0,
            lead=as_num(ws.cell(r, C["lead"]).value, 2.0) if C["lead"] else 2.0,
            mn=as_num(ws.cell(r, C["mn"]).value) if C["mn"] else 0.0,
            mx=as_num(ws.cell(r, C["mx"]).value) if C["mx"] else 0.0,
            price=as_num(ws.cell(r, C["price"]).value) if C["price"] else 0.0,
            receipt=as_dt(ws.cell(r, C["receipt"]).value) if C["receipt"] else None))
    return rows


def _loc(r):
    return " ".join(x for x in (r["box"], r["level"]) if x).strip()


def import_rows(s: Session, rows):
    """Write rows to the database, COMBINING rows that share a part code.

    The source workbooks list some parts on more than one row (same part stored
    in two boxes, or listed in both files). Physical stock is the SUM of those
    rows — previously the last row silently overwrote the others, so on-hand
    was wrong for those parts. Every merged part gets an opening-balance note
    listing each source row, so it can be verified by a physical count."""
    groups = {}
    for r in rows:
        groups.setdefault(r["code"], []).append(r)
    merged = []
    for code, g in groups.items():
        first = lambda k, d="": next((x[k] for x in g if x[k]), d)
        stock = round(sum(x["stock"] for x in g), 6)
        out_year = sum(x["out_year"] for x in g)
        lead = first("lead", 2.0) or 2.0
        mn = max(x["mn"] for x in g)
        mx = max(x["mx"] for x in g)
        rop = max(mn, round(out_year / 12.0 * lead, 1))
        locs = []
        for x in g:
            if _loc(x) and _loc(x) not in locs:
                locs.append(_loc(x))
        if len(locs) > 1:
            box, level = " + ".join(locs), ""
        else:
            box, level = first("box"), first("level")
        receipts = [x["receipt"] for x in g if x["receipt"]]
        it = upsert_item(
            s, code, category=cat_of(code), description=first("desc"),
            part_name=first("name"), part_number=first("pn"), brand=first("brand"),
            machine_group=first("machine"), uom=first("uom", "Pcs") or "Pcs",
            box=box, level=level, unit_price=first("price", 0.0) or 0.0, on_hand=stock,
            min_level=mn, max_level=max(mx, stock), reorder_point=rop, lead_time_months=lead,
            source_sheet=" + ".join(sorted({x["sheet"] for x in g})),
            last_receipt=max(receipts) if receipts else None, updated_at=datetime.utcnow())
        s.flush()
        if len(g) > 1:
            parts = "; ".join(f"{x['sheet'][:18]} แถว {x['row']} [{_loc(x) or '-'}] = {x['stock']:g}" for x in g)
            note = f"รวม {len(g)} แถวจากไฟล์ต้นฉบับ: {parts}"
            merged.append({"code": code, "rows": len(g), "total": stock, "detail": parts})
        else:
            note = "Opening balance from Excel"
        # opening balance ALWAYS recorded (also when 0) so the ledger explains on-hand
        s.add(StockTxn(item_id=it.id, item_code=code, txn_type="receive", qty=stock,
                       balance_after=stock, ref="OPENING", note=note[:1000]))
    s.commit()
    return {"items": len(groups), "rows": len(rows), "merged": merged}


def import_sources(s: Session, spare: Path = None, cons: Path = None, enrich=None):
    rows = []
    if spare and spare.exists():
        rows += read_sheet(spare, "Machine Spare parts 2026", enrich, kind="spare")
    if cons and cons.exists():
        rows += read_sheet(cons, "Tool And Facility", enrich, kind="tool")
    rep = import_rows(s, rows)
    print(f"  {rep['rows']} rows -> {rep['items']} parts ({len(rep['merged'])} parts combined from several rows)")
    return rep


def import_sheet(path: Path, sheet: str, s: Session, enrich=None, kind="spare"):
    """Backward-compatible single-sheet import."""
    return import_rows(s, read_sheet(path, sheet, enrich, kind))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=str(Path(__file__).resolve().parent.parent / "source"))
    args = ap.parse_args()
    d = Path(args.dir)
    f_spare = d / "2_Spare_Parts_Control_2026.xlsx"
    f_cons = d / "2_Spare_Parts_Control_2026_Consumable.xlsx"
    f_mach = d / "Machine_list_.xlsx"

    init_db()
    with SessionLocal() as s:
        print("Importing…")
        if f_mach.exists():
            import_machines(f_mach, s)
        enrich = enrich_from_itemcode(f_cons) if f_cons.exists() else {}
        import_sources(s, f_spare, f_cons, enrich)
        total = len(s.exec(select(Item)).all())
        print(f"Done. Total items in DB: {total}")


if __name__ == "__main__":
    main()
