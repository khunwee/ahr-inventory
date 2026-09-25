import io
import csv
import hmac
import base64
import hashlib
from datetime import datetime, date, timedelta
from contextlib import asynccontextmanager
from typing import Optional, List

from fastapi import (FastAPI, Depends, HTTPException, UploadFile, File, Form,
                     WebSocket, WebSocketDisconnect, Request, Query)
from fastapi.responses import StreamingResponse, FileResponse, JSONResponse, Response, Response
from fastapi.staticfiles import StaticFiles
from fastapi.security import OAuth2PasswordRequestForm
from pydantic import BaseModel, Field
from sqlalchemy import select, or_, func, update as sa_update

from .config import settings, UPLOAD_DIR, WEB_DIR, DATA_DIR
from .tz import local_now, local_today, to_local, utc_range, fmt_local
from .db import init_db, get_session, engine, urgency_of, SessionLocal, Session
from .models import (User, Machine, Item, Requisition, ReqLine, StockTxn,
                     Setting, PurchaseOrder, POLine, AuditLog, ReturnLog)
from .auth import hash_pw, verify_pw, make_token, current_user, require, ROLE_RANK
from .notifications import ws_manager, notify, test_channel
from .scheduler import start_scheduler
from . import settings_store, linebot, labels, images
from .audit import log as audit_log
from .services import (issue_requisition, confirm_message, reorder_message, change_stock, claim,
                       next_number, save_with_number)
import secrets
import httpx as _httpx


@asynccontextmanager
async def lifespan(app: FastAPI):
    init_db()
    seed_admin()
    start_scheduler()
    yield


app = FastAPI(title=settings.APP_NAME, lifespan=lifespan)


def seed_admin():
    with SessionLocal() as s:
        if not s.exec(select(User)).first():
            s.add(User(username="admin", full_name="System Admin",
                       password_hash=hash_pw("admin123"), role="admin",
                       must_change_pw=True))
            s.add(User(username="leader1", full_name="Line Leader",
                       password_hash=hash_pw("user123"), role="leader"))
            s.add(User(username="tech1", full_name="Technician 1",
                       password_hash=hash_pw("user123"), role="engineer"))
            s.commit()


# ----------------------------- schemas ------------------------------------
class LineIn(BaseModel):
    item_id: int
    qty: float = Field(gt=0, description="must be > 0")


class ReqIn(BaseModel):
    machine_id: Optional[int] = None
    machine_name: str = ""
    problem: str = ""
    note: str = ""
    lines: List[LineIn]


class CountIn(BaseModel):
    line_id: int
    counted_qty: float


class ReceiveIn(BaseModel):
    item_id: int
    qty: float = Field(gt=0, description="must be > 0")
    ref: str = ""
    note: str = ""


# ----------------------------- auth ---------------------------------------
@app.post("/api/auth/login")
def login(request: Request, form: OAuth2PasswordRequestForm = Depends(),
          s: Session = Depends(get_session)):
    u = s.exec(select(User).where(User.username == form.username)).first()
    if not u or not verify_pw(form.password, u.password_hash) or not u.active:
        raise HTTPException(401, "Wrong username or password")
    ip = request.client.host if request.client else ""
    audit_log("login", f"{u.username} signed in", user=u, ip=ip)
    tok = make_token(u)
    resp = JSONResponse({"access_token": tok, "token_type": "bearer",
                         "user": {"id": u.id, "name": u.full_name or u.username, "role": u.role,
                                  "must_change_pw": bool(u.must_change_pw)}})
    images.set_session_cookie(resp, request, tok)
    return resp


@app.get("/api/me")
def me(u: User = Depends(current_user)):
    return {"id": u.id, "username": u.username, "name": u.full_name or u.username,
            "role": u.role, "must_change_pw": bool(u.must_change_pw),
            "line_linked": bool(u.line_user_id)}


# --------------------------- items / search -------------------------------
def item_dict(it: Item):
    return {
        "id": it.id, "item_code": it.item_code, "description": it.description,
        "part_name": it.part_name, "part_number": it.part_number, "brand": it.brand,
        "machine_group": it.machine_group, "uom": it.uom, "box": it.box, "level": it.level,
        "on_hand": it.on_hand, "min_level": it.min_level, "reorder_point": it.reorder_point,
        "max_level": it.max_level, "lead_time_months": it.lead_time_months,
        "unit_price": it.unit_price, "category": it.category, "urgency": urgency_of(it),
        "image_ver": it.image_ver or 0,
    }


@app.get("/api/items")
def search_items(q: str = "", machine: str = "", low_only: bool = False,
                 limit: int = 60, offset: int = 0,
                 s: Session = Depends(get_session), u: User = Depends(current_user)):
    """Partial search across code, description, part name/number, brand, machine.
    Any fragment matches any field."""
    stmt = select(Item)
    for term in q.split():
        like = f"%{term}%"
        stmt = stmt.where(or_(
            Item.item_code.ilike(like), Item.description.ilike(like),
            Item.part_name.ilike(like), Item.part_number.ilike(like),
            Item.brand.ilike(like), Item.machine_group.ilike(like)))
    if machine:
        stmt = stmt.where(Item.machine_group.ilike(f"%{machine}%"))
    rows = s.exec(stmt.order_by(Item.item_code).limit(limit).offset(offset)).all()
    if low_only:
        rows = [r for r in rows if urgency_of(r)]
    return [item_dict(r) for r in rows]


@app.get("/api/items/{item_id}")
def get_item(item_id: int, s: Session = Depends(get_session), u: User = Depends(current_user)):
    it = s.get(Item, item_id)
    if not it:
        raise HTTPException(404, "Item not found")
    txns = s.exec(select(StockTxn).where(StockTxn.item_id == item_id)
                  .order_by(StockTxn.created_at.desc()).limit(20)).all()
    d = item_dict(it)
    d["history"] = [{"type": t.txn_type, "qty": t.qty, "balance_after": t.balance_after,
                     "ref": t.ref, "when": t.created_at, "machine": t.machine_name} for t in txns]
    return d


class ItemPatch(BaseModel):
    item_code: Optional[str] = None
    category: Optional[str] = None
    part_name: Optional[str] = None
    part_number: Optional[str] = None
    brand: Optional[str] = None
    description: Optional[str] = None
    machine_group: Optional[str] = None
    uom: Optional[str] = None
    box: Optional[str] = None
    level: Optional[str] = None
    unit_price: Optional[float] = None
    on_hand: Optional[float] = None
    min_level: Optional[float] = None
    reorder_point: Optional[float] = None
    max_level: Optional[float] = None
    lead_time_months: Optional[float] = None


@app.patch("/api/items/{item_id}")
def update_item(item_id: int, payload: ItemPatch, s: Session = Depends(get_session),
                u: User = Depends(require("admin"))):
    """Edit specific fields of one part. Changing on_hand writes a stock-adjust
    ledger entry so the balance stays auditable."""
    it = s.get(Item, item_id)
    if not it:
        raise HTTPException(404, "Item not found")
    data = payload.model_dump(exclude_unset=True)
    # item_code is the identity — allow change but keep it unique
    if "item_code" in data and data["item_code"] and data["item_code"] != it.item_code:
        code = data["item_code"].strip()
        if s.exec(select(Item).where(Item.item_code == code)).first():
            raise HTTPException(400, "รหัสอะไหล่นี้มีอยู่แล้ว")
        it.item_code = code
    for k, v in data.items():
        if k in ("on_hand", "item_code"):
            continue
        setattr(it, k, v)
    if "on_hand" in data and data["on_hand"] is not None:
        # admin typed an absolute counted value: lock the row, read the CURRENT
        # balance, and apply only the difference atomically (logged as adjust)
        cur = s.execute(select(Item.on_hand).where(Item.id == it.id).with_for_update()).scalar_one() or 0
        diff = round(float(data["on_hand"]) - cur, 6)
        if diff:
            bal = change_stock(s, it.id, diff)
            s.add(StockTxn(item_id=it.id, item_code=it.item_code, txn_type="adjust",
                           qty=diff, balance_after=bal, ref="EDIT",
                           note="แก้ไขจำนวนโดยแอดมิน", user_id=u.id))
    it.updated_at = datetime.utcnow()
    s.add(it); s.commit(); s.refresh(it)
    audit_log("item", f"edited {it.item_code}: {', '.join(data.keys())}", user=u)
    return {"ok": True, "item": item_dict(it)}


@app.get("/api/machines")
def machines(s: Session = Depends(get_session), u: User = Depends(current_user)):
    rows = s.exec(select(Machine).where(Machine.active == True).order_by(Machine.name)).all()
    return [{"id": m.id, "name": m.name} for m in rows]


# --------------------------- requisitions ---------------------------------
def _next_ref(s: Session) -> str:
    return next_number(s, Requisition.ref_no, f"REQ{local_now():%Y%m%d}-")


@app.post("/api/requisitions")
async def create_req(payload: ReqIn, s: Session = Depends(get_session),
                     u: User = Depends(current_user)):
    if not payload.lines:
        raise HTTPException(400, "Add at least one item to withdraw")
    mname = payload.machine_name
    if payload.machine_id:
        m = s.get(Machine, payload.machine_id)
        mname = m.name if m else mname
    req = Requisition(requester_id=u.id,
                      requester_name=u.full_name or u.username,
                      machine_id=payload.machine_id, machine_name=mname,
                      problem=payload.problem, note=payload.note, status="draft", source="web")
    req = save_with_number(s, req, "ref_no", _next_ref)
    s.commit(); s.refresh(req)
    for ln in payload.lines:
        it = s.get(Item, ln.item_id)
        if not it:
            continue
        s.add(ReqLine(requisition_id=req.id, item_id=it.id, item_code=it.item_code,
                      description=it.description, qty=ln.qty))
    s.commit()
    return {"id": req.id, "ref_no": req.ref_no, "status": req.status}


@app.post("/api/requisitions/{req_id}/confirm")
async def confirm_req(req_id: int, s: Session = Depends(get_session),
                      u: User = Depends(current_user)):
    """Confirm → deduct stock automatically, write ledger, fire notifications,
    check reorder thresholds."""
    req = s.get(Requisition, req_id)
    if not req:
        raise HTTPException(404, "Requisition not found")
    if req.status != "draft" or not claim(s, Requisition, req_id, "draft", "confirmed"):
        raise HTTPException(400, f"ใบเบิกนี้ถูกยืนยันไปแล้ว ({req.ref_no})")
    lines = s.exec(select(ReqLine).where(ReqLine.requisition_id == req_id)).all()
    triggered = []
    changed = []
    for ln in lines:
        it = s.get(Item, ln.item_id)
        if not it or not ln.qty or ln.qty <= 0:
            continue
        bal = change_stock(s, it.id, -ln.qty)              # atomic stock deduction
        ln.system_after = bal
        s.add(StockTxn(item_id=it.id, item_code=it.item_code, txn_type="issue",
                       qty=ln.qty, balance_after=bal, ref=req.ref_no,
                       user_id=u.id, machine_name=req.machine_name,
                       note=req.problem))
        s.add(ln)
        s.refresh(it)
        changed.append(item_dict(it))
        if urgency_of(it):
            triggered.append((it, urgency_of(it)))
    s.refresh(req)
    req.confirmed_at = datetime.utcnow()
    s.add(req); s.commit()
    audit_log("issue", f"{req.ref_no}: {len(lines)} line(s), machine={req.machine_name}", user=u)

    # real-time + chat notifications
    body = "\n".join(f"• {l.item_code}  x{l.qty:g}  ({l.description[:40]})" for l in lines)
    text = (f"📤 เบิกอะไหล่ {req.ref_no}\nผู้เบิก: {req.requester_name}\n"
            f"เครื่อง: {req.machine_name or '-'}\nปัญหา: {req.problem or '-'}\n{body}")
    await notify("requisition.confirmed", text,
                 {"ref": req.ref_no, "items": changed})
    # immediate alert for any part that just crossed a reorder threshold
    if triggered:
        atext = "🔔 ต้องสั่งซื้อเพิ่ม (หลังการเบิกล่าสุด):\n" + "\n".join(
            f"[{urg.upper()}] {it.item_code} คงเหลือ {it.on_hand:g} / ROP {it.reorder_point:g}"
            for it, urg in triggered)
        await notify("reorder.alert", atext,
                     {"items": [item_dict(it) for it, _ in triggered]})
    return {"status": "confirmed", "ref_no": req.ref_no,
            "reorder_triggered": [it.item_code for it, _ in triggered]}


@app.post("/api/requisitions/count")
def enter_count(payload: CountIn, s: Session = Depends(get_session),
                u: User = Depends(current_user)):
    """Technician enters the real physical remaining count after withdrawing,
    so a leader can reconcile it against system on-hand."""
    ln = s.get(ReqLine, payload.line_id)
    if not ln:
        raise HTTPException(404, "Line not found")
    ln.counted_qty = payload.counted_qty
    s.add(ln); s.commit()
    variance = None
    if ln.system_after is not None:
        variance = round(payload.counted_qty - ln.system_after, 3)
    return {"ok": True, "system_after": ln.system_after, "variance": variance}


@app.get("/api/requisitions")
def list_reqs(date_from: Optional[date] = None, date_to: Optional[date] = None,
              mine: bool = False, requester_id: Optional[int] = None,
              s: Session = Depends(get_session),
              u: User = Depends(current_user)):
    stmt = select(Requisition)
    if mine:
        stmt = stmt.where(Requisition.requester_id == u.id)
    if requester_id:
        stmt = stmt.where(Requisition.requester_id == requester_id)
    if date_from:
        stmt = stmt.where(Requisition.created_at >= utc_range(date_from, date_from)[0])
    if date_to:
        stmt = stmt.where(Requisition.created_at <= utc_range(date_to, date_to)[1])
    reqs = s.exec(stmt.order_by(Requisition.created_at.desc()).limit(300)).all()
    out = []
    for r in reqs:
        lines = s.exec(select(ReqLine).where(ReqLine.requisition_id == r.id)).all()
        rets = s.exec(select(ReturnLog).where(ReturnLog.requisition_id == r.id)
                      .order_by(ReturnLog.created_at)).all()
        out.append({
            "id": r.id, "ref_no": r.ref_no, "requester": r.requester_name,
            "machine": r.machine_name, "problem": r.problem, "status": r.status,
            "source": r.source, "created_at": r.created_at, "photo": r.photo_path,
            "has_return": bool(rets),
            "remark": "; ".join(f"คืน {x.item_code} x{x.qty:g} — {x.reason}" for x in rets),
            "returns": [{"item_code": x.item_code, "qty": x.qty, "reason": x.reason,
                         "by": x.returned_by, "at": x.created_at} for x in rets],
            "lines": [{"line_id": l.id, "item_id": l.item_id, "item_code": l.item_code,
                       "image_ver": (getattr(s.get(Item, l.item_id), "image_ver", 0) or 0),
                       "description": l.description, "qty": l.qty,
                       "returned_qty": l.returned_qty or 0,
                       "counted_qty": l.counted_qty, "system_after": l.system_after,
                       "variance": (None if l.counted_qty is None or l.system_after is None
                                    else round(l.counted_qty - l.system_after, 3))}
                      for l in lines]})
    return out


@app.get("/api/requesters")
def requesters(s: Session = Depends(get_session), u: User = Depends(current_user)):
    """People who can/do withdraw (for the technician filter) with their counts."""
    counts = {}
    for r in s.exec(select(Requisition)).all():
        if r.requester_id:
            counts[r.requester_id] = counts.get(r.requester_id, 0) + 1
    users = s.exec(select(User).where(User.active == True).order_by(User.username)).all()
    out = [{"id": x.id, "name": x.full_name or x.username, "role": x.role,
            "count": counts.get(x.id, 0)} for x in users]
    # keep engineers/leaders/admins (people who withdraw); viewers excluded
    out = [o for o in out if o["role"] != "viewer"]
    out.sort(key=lambda o: (-o["count"], o["name"]))
    return out


# ------------------------------ receive -----------------------------------
@app.post("/api/receive")
async def receive(payload: ReceiveIn, s: Session = Depends(get_session),
                  u: User = Depends(require("leader"))):
    it = s.get(Item, payload.item_id)
    if not it:
        raise HTTPException(404, "Item not found")
    bal = change_stock(s, it.id, payload.qty)            # atomic
    it.last_receipt = datetime.utcnow()
    s.add(StockTxn(item_id=it.id, item_code=it.item_code, txn_type="receive",
                   qty=payload.qty, balance_after=bal, ref=payload.ref,
                   user_id=u.id, note=payload.note))
    s.add(it); s.commit(); s.refresh(it)
    audit_log("receive", f"{it.item_code} +{payload.qty:g} -> {it.on_hand:g}", user=u)
    await notify("stock.received",
                 f"📥 รับเข้า {it.item_code} x{payload.qty:g} คงเหลือ {it.on_hand:g}",
                 {"item": item_dict(it)}, channels=("ws",))
    return {"ok": True, "on_hand": it.on_hand}


# ------------------------------ reorder -----------------------------------
@app.get("/api/reorder")
def reorder(s: Session = Depends(get_session), u: User = Depends(current_user)):
    items = s.exec(select(Item)).all()
    rank = {"critical": 0, "high": 1, "watch": 2}
    flagged = [(it, urgency_of(it)) for it in items if urgency_of(it)]
    flagged.sort(key=lambda x: (rank.get(x[1], 9), -(x[0].reorder_point - x[0].on_hand)))
    return [{**item_dict(it), "urgency": urg,
             "suggest_qty": max(round((it.max_level or it.reorder_point * 2) - it.on_hand, 2), 0)}
            for it, urg in flagged]


# ----------------------------- dashboard ----------------------------------
@app.get("/api/dashboard")
def dashboard(s: Session = Depends(get_session), u: User = Depends(current_user)):
    items = s.exec(select(Item)).all()
    total_items = len(items)
    total_value = sum((it.on_hand or 0) * (it.unit_price or 0) for it in items)
    urgencies = [urgency_of(it) for it in items]
    crit = urgencies.count("critical")
    high = urgencies.count("high")
    watch = urgencies.count("watch")
    out_of_stock = sum(1 for it in items if (it.on_hand or 0) <= 0)

    # issues over last 14 days
    since = datetime.utcnow() - timedelta(days=14)
    txns = s.exec(select(StockTxn).where(StockTxn.created_at >= since)).all()
    by_day = {}
    for t in txns:
        d = to_local(t.created_at).strftime("%m-%d")
        by_day.setdefault(d, {"issue": 0, "receive": 0})
        if t.txn_type in ("issue", "receive"):
            by_day[d][t.txn_type] += t.qty
    days = sorted(by_day.keys())

    # top consumed parts (30d)
    since30 = datetime.utcnow() - timedelta(days=30)
    consumed = {}
    for t in txns:
        if t.txn_type == "issue" and t.created_at >= since30:
            consumed[t.item_code] = consumed.get(t.item_code, 0) + t.qty
    top = sorted(consumed.items(), key=lambda x: -x[1])[:8]

    # value by category
    by_cat = {}
    for it in items:
        by_cat[it.category or "?"] = by_cat.get(it.category or "?", 0) + (it.on_hand or 0) * (it.unit_price or 0)

    # --- executive metrics ---
    # reorder status breakdown
    ok = total_items - crit - high - watch
    reorder_breakdown = {"labels": ["ปกติ", "เฝ้าระวัง", "ใกล้หมด", "ต้องสั่งด่วน"],
                         "values": [ok, watch, high, crit]}

    # top machines by withdrawn qty (30d)
    all_issue = s.exec(select(StockTxn).where(StockTxn.txn_type == "issue")).all()
    price_map = {it.id: (it.unit_price or 0) for it in items}
    since30b = datetime.utcnow() - timedelta(days=30)
    by_machine = {}
    for t in all_issue:
        if t.created_at >= since30b:
            m = (t.machine_name or "ไม่ระบุ").strip() or "ไม่ระบุ"
            by_machine[m] = by_machine.get(m, 0) + t.qty
    tm = sorted(by_machine.items(), key=lambda x: -x[1])[:8]

    # withdrawal value by month (last 6 months)
    monthly = {}
    for t in all_issue:
        key = to_local(t.created_at).strftime("%Y-%m")
        monthly[key] = monthly.get(key, 0) + t.qty * price_map.get(t.item_id, 0)
    mkeys = sorted(monthly.keys())[-6:]

    # PO status
    pos = s.exec(select(PurchaseOrder)).all()
    po_status = {"ordered": sum(1 for p in pos if p.status == "ordered"),
                 "received": sum(1 for p in pos if p.status == "received"),
                 "cancelled": sum(1 for p in pos if p.status == "cancelled")}

    return {
        "total_items": total_items, "total_value": round(total_value, 2),
        "critical": crit, "high": high, "watch": watch, "out_of_stock": out_of_stock,
        "trend": {"days": days,
                  "issue": [round(by_day[d]["issue"], 1) for d in days],
                  "receive": [round(by_day[d]["receive"], 1) for d in days]},
        "top_consumed": {"labels": [c for c, _ in top], "values": [round(v, 1) for _, v in top]},
        "by_category": {"labels": list(by_cat.keys()),
                        "values": [round(v, 0) for v in by_cat.values()]},
        "reorder_breakdown": reorder_breakdown,
        "by_machine": {"labels": [m for m, _ in tm], "values": [round(v, 1) for _, v in tm]},
        "monthly_value": {"labels": mkeys, "values": [round(monthly[k], 0) for k in mkeys]},
        "po_status": po_status,
    }


# --------------------------- oracle export --------------------------------
def _issue_lines(s, date_from, date_to, requester_id=None):
    """Every confirmed withdrawal line in the Thai-time date range, with returns
    netted out: yields (requisition, line, item, unit_price, returned, net_qty)."""
    lo, hi = utc_range(date_from, date_to)
    stmt = select(Requisition).where(
        Requisition.status.in_(["confirmed", "reconciled"]),
        Requisition.created_at >= lo, Requisition.created_at <= hi)
    if requester_id:
        stmt = stmt.where(Requisition.requester_id == requester_id)
    for r in s.exec(stmt.order_by(Requisition.created_at)).all():
        for l in s.exec(select(ReqLine).where(ReqLine.requisition_id == r.id)).all():
            it = s.get(Item, l.item_id)
            price = (it.unit_price if it else 0) or 0
            ret = l.returned_qty or 0
            yield r, l, it, price, ret, round(l.qty - ret, 3)


ORACLE_HEADER = ["RefNo", "Date", "ItemCode", "Description", "Qty", "Returned", "NetQty",
                 "UOM", "Machine", "Problem", "Requester", "UnitPrice", "Amount", "Remark"]


def _return_remarks(s, reqline_id):
    logs = s.exec(select(ReturnLog).where(ReturnLog.reqline_id == reqline_id)).all()
    return "; ".join(f"คืน {x.qty:g}: {x.reason}" for x in logs)


def _oracle_rows(s, date_from, date_to, requester_id=None):
    rows = []
    for r, l, it, price, ret, net in _issue_lines(s, date_from, date_to, requester_id):
        rows.append([r.ref_no, fmt_local(r.created_at, "%d/%m/%Y %H:%M"), l.item_code,
                     l.description, l.qty, ret, net, it.uom if it else "", r.machine_name,
                     r.problem, r.requester_name, price, round(net * price, 2),
                     _return_remarks(s, l.id) if ret else ""])
    return rows


@app.get("/api/export/oracle")
def export_oracle(date_from: date, date_to: date, requester_id: Optional[int] = None,
                  s: Session = Depends(get_session), u: User = Depends(require("leader"))):
    """CSV of confirmed issues (net of returns), ready to key into Oracle."""
    buf = io.StringIO()
    buf.write("\ufeff")                       # BOM so Excel shows Thai correctly
    w = csv.writer(buf)
    w.writerow(ORACLE_HEADER)
    for row in _oracle_rows(s, date_from, date_to, requester_id):
        w.writerow(row)
    fname = f"oracle_export_{date_from}_{date_to}.csv"
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv; charset=utf-8",
                             headers={"Content-Disposition": f"attachment; filename={fname}"})


@app.get("/api/uploads/{name}")
def get_upload(name: str, u: User = Depends(images.viewer)):
    if "/" in name or "\\" in name or ".." in name:
        raise HTTPException(400, "bad name")
    p = UPLOAD_DIR / name
    if not p.exists():
        raise HTTPException(404, "Not found")
    return FileResponse(p)


# ------------------------------ websocket ---------------------------------
@app.websocket("/ws")
async def ws_endpoint(ws: WebSocket):
    await ws_manager.connect(ws)
    try:
        while True:
            await ws.receive_text()   # keepalive / ignore inbound
    except WebSocketDisconnect:
        ws_manager.disconnect(ws)


# --------------------------- LINE webhook ---------------------------------
# --------------------------- LINE webhook ---------------------------------
async def _line_reply(reply_token: str, messages: list):
    token = settings_store.get("LINE_CHANNEL_ACCESS_TOKEN")
    if not token or not reply_token:
        return
    async with _httpx.AsyncClient(timeout=10) as c:
        await c.post("https://api.line.me/v2/bot/message/reply",
                     headers={"Authorization": f"Bearer {token}"},
                     json={"replyToken": reply_token, "messages": messages[:5]})


async def _line_get_image(message_id: str) -> bytes:
    token = settings_store.get("LINE_CHANNEL_ACCESS_TOKEN")
    async with _httpx.AsyncClient(timeout=15) as c:
        r = await c.get(f"https://api-data.line.me/v2/bot/message/{message_id}/content",
                        headers={"Authorization": f"Bearer {token}"})
        return r.content


@app.post("/line/webhook")
async def line_webhook(request: Request):
    """LINE Messaging API webhook — drives the requisition bot (see linebot.py)."""
    body = await request.body()
    secret = settings_store.get("LINE_CHANNEL_SECRET")
    if secret:
        mac = hmac.new(secret.encode(), body, hashlib.sha256).digest()
        if not hmac.compare_digest(base64.b64encode(mac).decode(),
                                   request.headers.get("x-line-signature", "")):
            return JSONResponse({"ok": False, "reason": "bad signature"}, status_code=403)
    import json as _json
    data = _json.loads(body or b"{}")
    for ev in data.get("events", []):
        if ev.get("type") != "message":
            continue
        src = ev.get("source", {})
        line_uid = src.get("userId", "")
        reply_token = ev.get("replyToken", "")
        msg = ev.get("message", {})
        if msg.get("type") == "text":
            result = linebot.handle_text(line_uid, msg.get("text", ""))
            confirmed = None
            if isinstance(result, tuple):
                messages, confirmed = result
            else:
                messages = result
            await _line_reply(reply_token, messages)
            if confirmed:
                req, lines, triggered = confirmed
                await notify("requisition.confirmed", confirm_message(req, lines),
                             {"ref": req.ref_no})
                if triggered:
                    await notify("reorder.alert", reorder_message(triggered),
                                 {"items": [item_dict(it) for it, _ in triggered]})
        elif msg.get("type") == "image":
            try:
                content = await _line_get_image(msg.get("id", ""))
                rid = linebot.last_req_id(line_uid)
                ok = False
                if rid and content:
                    with SessionLocal() as s2:
                        images.save_req_photo_bytes(s2, rid, content); s2.commit(); ok = True
                await _line_reply(reply_token, [{"type": "text",
                    "text": "📎 แนบรูปกับใบเบิกล่าสุดแล้ว" if ok else "ยังไม่มีใบเบิกให้แนบรูป"}])
            except Exception:
                pass
    return {"ok": True}


@app.post("/api/me/line-code")
def line_link_code(u: User = Depends(current_user)):
    """Generate a one-time code to link this account to LINE."""
    code = f"{secrets.randbelow(1000000):06d}"
    linebot.new_link_code(u.id, code)
    return {"code": code, "expires_min": 10}


# --------------------------- user management ------------------------------
class UserIn(BaseModel):
    username: str
    full_name: str = ""
    role: str = "engineer"
    password: str


class UserPatch(BaseModel):
    full_name: Optional[str] = None
    role: Optional[str] = None
    active: Optional[bool] = None
    password: Optional[str] = None


class PwdIn(BaseModel):
    old_password: str
    new_password: str


@app.get("/api/users")
def list_users(s: Session = Depends(get_session), u: User = Depends(require("admin"))):
    rows = s.exec(select(User).order_by(User.username)).all()
    return [{"id": x.id, "username": x.username, "full_name": x.full_name,
             "role": x.role, "active": x.active} for x in rows]


@app.post("/api/users")
def create_user(payload: UserIn, s: Session = Depends(get_session),
                u: User = Depends(require("admin"))):
    if payload.role not in ROLE_RANK:
        raise HTTPException(400, "Invalid role")
    if s.exec(select(User).where(User.username == payload.username.strip())).first():
        raise HTTPException(400, "Username already exists")
    if len(payload.password) < 4:
        raise HTTPException(400, "Password must be at least 4 characters")
    nu = User(username=payload.username.strip(), full_name=payload.full_name,
              role=payload.role, password_hash=hash_pw(payload.password),
              must_change_pw=True)
    s.add(nu); s.commit()
    audit_log("user", f"created {nu.username} ({nu.role})", user=u)
    return {"ok": True, "id": nu.id}


@app.patch("/api/users/{uid}")
def update_user(uid: int, payload: UserPatch, s: Session = Depends(get_session),
                u: User = Depends(require("admin"))):
    x = s.get(User, uid)
    if not x:
        raise HTTPException(404, "User not found")
    if x.id == u.id and payload.active is False:
        raise HTTPException(400, "You can't deactivate your own account")
    if payload.full_name is not None:
        x.full_name = payload.full_name
    if payload.role is not None:
        if payload.role not in ROLE_RANK:
            raise HTTPException(400, "Invalid role")
        x.role = payload.role
    if payload.active is not None:
        x.active = payload.active
    if payload.password:
        if len(payload.password) < 4:
            raise HTTPException(400, "Password must be at least 4 characters")
        x.password_hash = hash_pw(payload.password)
    s.add(x); s.commit()
    return {"ok": True}


@app.post("/api/me/password")
def change_my_password(payload: PwdIn, s: Session = Depends(get_session),
                       u: User = Depends(current_user)):
    if not verify_pw(payload.old_password, u.password_hash):
        raise HTTPException(400, "Current password is incorrect")
    if len(payload.new_password) < 4:
        raise HTTPException(400, "New password must be at least 4 characters")
    u.password_hash = hash_pw(payload.new_password)
    u.must_change_pw = False
    s.add(u); s.commit()
    return {"ok": True}


# --------------------------- settings (admin) -----------------------------
class SettingsIn(BaseModel):
    values: dict


@app.get("/api/settings")
def get_settings(u: User = Depends(require("admin"))):
    return settings_store.masked()


@app.put("/api/settings")
def put_settings(payload: SettingsIn, u: User = Depends(require("admin"))):
    # ignore empty secret values so we don't wipe a stored secret with a blank
    clean = {k: v for k, v in payload.values.items()
             if not (k in settings_store.SECRET_KEYS and (v is None or v == ""))}
    settings_store.set_many(clean)
    audit_log("setting", f"updated {list(clean.keys())}", user=u)
    return {"ok": True}


@app.post("/api/settings/test/{channel}")
async def settings_test(channel: str, u: User = Depends(require("admin"))):
    return await test_channel(channel)


# ------------------------------ audit log ---------------------------------
@app.get("/api/audit")
def get_audit(limit: int = 200, s: Session = Depends(get_session),
              u: User = Depends(require("leader"))):
    rows = s.exec(select(AuditLog).order_by(AuditLog.ts.desc()).limit(limit)).all()
    return [{"ts": r.ts, "username": r.username, "role": r.role,
             "action": r.action, "detail": r.detail, "ip": r.ip} for r in rows]


# ------------------------------ purchase orders ---------------------------
class POItemIn(BaseModel):
    item_id: int
    qty: float


class POIn(BaseModel):
    note: str = ""
    lines: List[POItemIn]


def _next_po(s) -> str:
    return next_number(s, PurchaseOrder.po_no, f"PO{local_now():%Y%m}-")


@app.post("/api/po")
def create_po(payload: POIn, s: Session = Depends(get_session),
              u: User = Depends(require("leader"))):
    if not payload.lines:
        raise HTTPException(400, "No items to order")
    po = PurchaseOrder(created_by=u.full_name or u.username,
                       note=payload.note, status="ordered")
    po = save_with_number(s, po, "po_no", _next_po)
    for ln in payload.lines:
        it = s.get(Item, ln.item_id)
        if not it:
            continue
        s.add(POLine(po_id=po.id, item_id=it.id, item_code=it.item_code,
                     description=it.part_name or it.description, qty=ln.qty,
                     unit_price=it.unit_price))
    s.commit()
    audit_log("po", f"created {po.po_no} ({len(payload.lines)} lines)", user=u)
    return {"ok": True, "id": po.id, "po_no": po.po_no}


@app.get("/api/po")
def list_po(s: Session = Depends(get_session), u: User = Depends(require("leader"))):
    out = []
    for po in s.exec(select(PurchaseOrder).order_by(PurchaseOrder.created_at.desc()).limit(200)).all():
        lines = s.exec(select(POLine).where(POLine.po_id == po.id)).all()
        out.append({"id": po.id, "po_no": po.po_no, "status": po.status,
                    "created_by": po.created_by, "created_at": po.created_at,
                    "note": po.note, "lines": [
                        {"id": l.id, "item_code": l.item_code, "description": l.description,
                         "qty": l.qty, "received_qty": l.received_qty,
                         "unit_price": l.unit_price} for l in lines]})
    return out


@app.post("/api/po/{po_id}/receive")
async def receive_po(po_id: int, s: Session = Depends(get_session),
                     u: User = Depends(require("leader"))):
    """Mark a PO received → add every line's qty into stock (ledger + on_hand)."""
    po = s.get(PurchaseOrder, po_id)
    if not po:
        raise HTTPException(404, "PO not found")
    if po.status != "ordered" or not claim(s, PurchaseOrder, po_id, "ordered", "received"):
        raise HTTPException(400, "ใบสั่งซื้อนี้รับของไปแล้ว หรือถูกยกเลิกแล้ว — รับซ้ำไม่ได้"
                                 if po.status != "cancelled" else "ใบสั่งซื้อถูกยกเลิกแล้ว — รับของไม่ได้")
    lines = s.exec(select(POLine).where(POLine.po_id == po_id)).all()
    for l in lines:
        it = s.get(Item, l.item_id)
        if not it or not l.qty or l.qty <= 0:
            continue
        bal = change_stock(s, it.id, l.qty)               # atomic
        it.last_receipt = datetime.utcnow()
        l.received_qty = l.qty
        s.add(StockTxn(item_id=it.id, item_code=it.item_code, txn_type="receive",
                       qty=l.qty, balance_after=bal, ref=po.po_no,
                       user_id=u.id, note="PO receive"))
        s.add(it); s.add(l)
    s.refresh(po)
    po.received_at = datetime.utcnow()
    s.add(po); s.commit()
    audit_log("po", f"received {po.po_no}", user=u)
    await notify("stock.received", f"📥 รับของตาม {po.po_no} เข้าสต็อกแล้ว",
                 {"po": po.po_no}, channels=("ws",))
    return {"ok": True, "status": "received"}


@app.post("/api/po/{po_id}/cancel")
def cancel_po(po_id: int, s: Session = Depends(get_session),
              u: User = Depends(require("leader"))):
    po = s.get(PurchaseOrder, po_id)
    if not po:
        raise HTTPException(404, "PO not found")
    if po.status == "received":
        raise HTTPException(400, "Cannot cancel a received PO")
    po.status = "cancelled"
    s.add(po); s.commit()
    return {"ok": True}


@app.get("/api/po/{po_id}/export.xlsx")
def po_xlsx(po_id: int, s: Session = Depends(get_session),
            u: User = Depends(require("leader"))):
    po = s.get(PurchaseOrder, po_id)
    if not po:
        raise HTTPException(404, "PO not found")
    lines = s.exec(select(POLine).where(POLine.po_id == po_id)).all()
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "PO"
    ws.append([f"ใบสั่งซื้อ {po.po_no}"])
    ws.append(["วันที่", fmt_local(po.created_at, "%d/%m/%Y %H:%M"), "โดย", po.created_by])
    ws.append([])
    ws.append(["No.", "Item Code", "Description", "Qty", "Unit Price", "Amount"])
    for i, l in enumerate(lines, 1):
        ws.append([i, l.item_code, l.description, l.qty, l.unit_price,
                   round(l.qty * l.unit_price, 2)])
    ws.append([])
    ws.append(["", "", "", "", "รวม", round(sum(l.qty * l.unit_price for l in lines), 2)])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={po.po_no}.xlsx"})


# --------------------------- oracle export (xlsx) -------------------------
@app.get("/api/export/oracle.xlsx")
def export_oracle_xlsx(date_from: date, date_to: date, requester_id: Optional[int] = None,
                       s: Session = Depends(get_session), u: User = Depends(require("leader"))):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Issues"
    ws.append(ORACLE_HEADER)
    for row in _oracle_rows(s, date_from, date_to, requester_id):
        ws.append(row)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"oracle_export_{date_from}_{date_to}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={fname}"})


# ------------------------------ QR labels ---------------------------------
@app.get("/api/labels/{item_id}.svg")
def label_svg(item_id: int, s: Session = Depends(get_session)):
    it = s.get(Item, item_id)
    if not it:
        raise HTTPException(404, "Item not found")
    return Response(content=labels.qr_svg(it.item_code), media_type="image/svg+xml")


@app.get("/api/items/by-code/{code}")
def item_by_code(code: str, s: Session = Depends(get_session), u: User = Depends(current_user)):
    it = s.exec(select(Item).where(Item.item_code == code.strip())).first()
    if not it:
        raise HTTPException(404, "Item not found")
    return item_dict(it)


# --------------------------- dead stock / aging ---------------------------
@app.get("/api/deadstock")
def deadstock(s: Session = Depends(get_session), u: User = Depends(current_user)):
    now = datetime.utcnow()
    items = s.exec(select(Item)).all()
    buckets = {"0-90": 0, "91-180": 0, "181-365": 0, ">365": 0, "no-move": 0}
    dead = []
    for it in items:
        if (it.on_hand or 0) <= 0:
            continue
        ref = it.last_receipt or it.updated_at
        days = (now - ref).days if ref else None
        val = (it.on_hand or 0) * (it.unit_price or 0)
        if days is None:
            buckets["no-move"] += val
        elif days <= 90:
            buckets["0-90"] += val
        elif days <= 180:
            buckets["91-180"] += val
        elif days <= 365:
            buckets["181-365"] += val
        else:
            buckets[">365"] += val
        if days is not None and days > settings.DEAD_STOCK_DAYS:
            dead.append({**item_dict(it), "age_days": days, "stock_value": round(val, 2)})
    dead.sort(key=lambda x: -x["stock_value"])
    return {"buckets": {k: round(v, 0) for k, v in buckets.items()},
            "dead_count": len(dead), "dead_value": round(sum(d["stock_value"] for d in dead), 2),
            "items": dead[:100]}


# --------------------------- optimization ---------------------------------
class OptIn(BaseModel):
    mode: str = "max_coverage"
    budget: Optional[float] = None
    only_flagged: bool = True


@app.post("/api/optimize/reorder")
def optimize_reorder_api(payload: OptIn, u: User = Depends(require("leader"))):
    from . import optimize as opt
    result = opt.optimize(mode=payload.mode, budget=payload.budget,
                          only_flagged=payload.only_flagged)
    audit_log("optimize", f"mode={payload.mode} method={result.get('method')} "
                          f"ordered={result.get('items_ordered')}", user=u)
    return result


@app.get("/api/optimize/status")
def optimize_status(u: User = Depends(require("leader"))):
    from . import optimize as opt
    return opt.solver_status()


@app.post("/api/optimize/apply-minmax")
def optimize_apply_minmax(u: User = Depends(require("leader"))):
    from . import optimize as opt
    r = opt.apply_minmax()
    audit_log("optimize", f"applied min/max to {r.get('updated')} items", user=u)
    return r


@app.get("/api/export/stock.xlsx")
def export_stock(q: str = "", low_only: bool = False,
                 s: Session = Depends(get_session), u: User = Depends(current_user)):
    """Full stock database as Excel (respects the current search/low filter)."""
    stmt = select(Item)
    for term in q.split():
        like = f"%{term}%"
        stmt = stmt.where(or_(Item.item_code.ilike(like), Item.description.ilike(like),
                              Item.part_name.ilike(like), Item.part_number.ilike(like),
                              Item.brand.ilike(like), Item.machine_group.ilike(like)))
    items = s.exec(stmt.order_by(Item.item_code)).all()
    if low_only:
        items = [it for it in items if urgency_of(it)]
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Stock"
    ws.append(["Item Code", "Category", "Part Name", "Part Number", "Brand", "Machine",
               "UOM", "Box", "Level", "On Hand", "Min", "Reorder Point", "Max",
               "Unit Price", "Stock Value", "Status"])
    for it in items:
        ws.append([it.item_code, it.category, it.part_name, it.part_number, it.brand,
                   it.machine_group, it.uom, it.box, it.level, it.on_hand, it.min_level,
                   it.reorder_point, it.max_level, it.unit_price,
                   round((it.on_hand or 0) * (it.unit_price or 0), 2),
                   {"critical": "CRITICAL", "high": "LOW", "watch": "WATCH"}.get(urgency_of(it), "OK")])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=stock_database.xlsx"})


@app.get("/api/requisitions/summary")
def requisitions_summary(date_from: date, date_to: date, requester_id: Optional[int] = None,
                         s: Session = Depends(get_session), u: User = Depends(require("leader"))):
    """Aggregated withdrawal summary (net of returns) — reference for Oracle."""
    agg, reqs, total_value = {}, set(), 0.0
    for r, l, it, price, ret, net in _issue_lines(s, date_from, date_to, requester_id):
        reqs.add(r.id)
        a = agg.setdefault(l.item_code, {"item_code": l.item_code, "description": l.description,
                                         "uom": it.uom if it else "", "qty": 0.0, "returned": 0.0,
                                         "unit_price": price, "machines": set()})
        a["qty"] += l.qty; a["returned"] += ret
        if r.machine_name:
            a["machines"].add(r.machine_name)
        total_value += net * price
    rows = [{"item_code": v["item_code"], "description": v["description"], "uom": v["uom"],
             "qty": round(v["qty"] - v["returned"], 2), "gross": round(v["qty"], 2),
             "returned": round(v["returned"], 2), "unit_price": v["unit_price"],
             "amount": round((v["qty"] - v["returned"]) * v["unit_price"], 2),
             "machines": ", ".join(sorted(v["machines"]))[:60]}
            for v in agg.values()]
    rows.sort(key=lambda x: -x["amount"])
    return {"date_from": str(date_from), "date_to": str(date_to), "requisitions": len(reqs),
            "items": len(rows), "total_value": round(total_value, 2), "rows": rows}


# --------------------------- data import (admin) --------------------------
IMPORT_COLS = ["item_code", "category", "description", "part_name", "part_number",
               "brand", "machine_group", "uom", "box", "level", "unit_price",
               "on_hand", "min_level", "reorder_point", "lead_time_months"]


@app.get("/api/import/template.xlsx")
def import_template(u: User = Depends(require("admin"))):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active; ws.title = "Items"
    ws.append(IMPORT_COLS)
    ws.append(["MSP0000001A", "MSP", "ตัวอย่าง: PUMP \\ GEAR PUMP : GP-123 \\ BRAND",
               "GEAR PUMP", "GP-123", "BRAND", "Komatsu 800T.", "Pcs", "A1", "1",
               1500, 5, 2, 3, 2])
    ws2 = wb.create_sheet("Machines")
    ws2.append(["name"])
    ws2.append(["Komatsu 800T."])
    ws2.append(["Komatsu 1200T."])
    # help sheet
    h = wb.create_sheet("README")
    for line in [["วิธีใช้ Template นำเข้าข้อมูล"],
                 ["1) กรอกอะไหล่ในชีต 'Items' (ต้องมี item_code ไม่ซ้ำ)"],
                 ["2) กรอกรายชื่อเครื่องในชีต 'Machines'"],
                 ["3) อัปโหลดไฟล์นี้ในหน้า 'นำเข้าข้อมูล' ของโปรแกรม"],
                 ["หมายเหตุ: on_hand = จำนวนคงเหลือเริ่มต้น, reorder_point = จุดสั่งซื้อ"]]:
        h.append(line)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=import_template.xlsx"})


@app.post("/api/import/upload")
async def import_upload(file: UploadFile = File(...), replace: bool = False,
                        s: Session = Depends(get_session), u: User = Depends(require("admin"))):
    """Import items + machines from a filled template. Upserts by item_code.
    If replace=True, wipes existing items/machines/stock/requisitions/POs first
    so the site can start fresh on a new dataset."""
    from openpyxl import load_workbook
    data = await file.read()
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        raise HTTPException(400, "ไฟล์ไม่ถูกต้อง (ต้องเป็น .xlsx)")

    if replace:
        for m in (POLine, PurchaseOrder, ReqLine, Requisition, StockTxn, Item, Machine):
            s.query(m).delete()
        s.commit()

    def num(v, d=0.0):
        try:
            return float(v)
        except (TypeError, ValueError):
            return d

    n_items = n_mach = 0
    merged = []
    if "Items" in wb.sheetnames:
        ws = wb["Items"]
        header = [str(c.value).strip() if c.value else "" for c in ws[1]]
        idx = {name: header.index(name) for name in IMPORT_COLS if name in header}
        if "item_code" not in idx:
            raise HTTPException(400, "ไม่พบคอลัมน์ item_code ในชีต Items")
        # 1) read + group rows by part code (duplicates are COMBINED: stock summed)
        groups = {}
        for rn, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or idx["item_code"] >= len(row):
                continue
            code = row[idx["item_code"]]
            if not code or not str(code).strip():
                continue
            groups.setdefault(str(code).strip(), []).append((rn, row))
        # 2) write one record per code
        for code, grp in groups.items():
            rn, row = grp[0]
            g = lambda k, d="", row=row: (str(row[idx[k]]).strip() if k in idx and idx[k] < len(row) and row[idx[k]] is not None else d)
            gn = lambda k, d=0.0, r=None: (num(r[idx[k]], d) if k in idx and idx[k] < len(r) else d)
            fields = dict(category=g("category"), description=g("description"),
                          part_name=g("part_name"), part_number=g("part_number"),
                          brand=g("brand"), machine_group=g("machine_group"),
                          uom=g("uom", "Pcs") or "Pcs", box=g("box"), level=g("level"),
                          unit_price=gn("unit_price", 0.0, row), min_level=gn("min_level", 0.0, row),
                          reorder_point=gn("reorder_point", 0.0, row),
                          lead_time_months=gn("lead_time_months", 2.0, row),
                          updated_at=datetime.utcnow())
            oh = round(sum(gn("on_hand", 0.0, r) for _, r in grp), 6)
            note = "Imported/updated via template"
            if len(grp) > 1:
                merged.append(code)
                note = f"รวม {len(grp)} แถวใน template (แถว {', '.join(str(x[0]) for x in grp)})"
            it = s.exec(select(Item).where(Item.item_code == code)).first()
            if it:
                for k, v in fields.items():
                    setattr(it, k, v)
                it.on_hand = oh
            else:
                it = Item(item_code=code, on_hand=oh, **fields)
                s.add(it)
            s.flush()
            s.add(StockTxn(item_id=it.id, item_code=code, txn_type="adjust",
                           qty=oh, balance_after=oh, ref="IMPORT", note=note, user_id=u.id))
            n_items += 1
    if "Machines" in wb.sheetnames:
        ws = wb["Machines"]
        existing = {m.name for m in s.exec(select(Machine)).all()}
        for row in ws.iter_rows(min_row=2, values_only=True):
            name = row[0] if row else None
            if name and str(name).strip() and str(name).strip() not in existing:
                s.add(Machine(name=str(name).strip())); existing.add(str(name).strip()); n_mach += 1
    s.commit()
    audit_log("import", f"items={n_items}, machines={n_mach}, merged={len(merged)}", user=u)
    return {"ok": True, "items": n_items, "machines": n_mach, "merged": merged}


# --------------------------- public config --------------------------------
@app.get("/api/config")
def public_config():
    """Non-secret settings the browser needs (time zone for date display)."""
    return {"timezone": settings.TIMEZONE, "app": settings.APP_NAME,
            "server_time_local": local_now().strftime("%d/%m/%Y %H:%M")}


# --------------------------- create new item (admin) ----------------------
class ItemCreate(BaseModel):
    item_code: str
    category: str = ""
    part_name: str = ""
    part_number: str = ""
    brand: str = ""
    description: str = ""
    machine_group: str = ""
    uom: str = "Pcs"
    box: str = ""
    level: str = ""
    unit_price: float = 0.0
    on_hand: float = 0.0
    min_level: float = 0.0
    reorder_point: float = 0.0
    max_level: float = 0.0
    lead_time_months: float = 2.0


@app.post("/api/items")
def create_item(payload: ItemCreate, s: Session = Depends(get_session),
                u: User = Depends(require("admin"))):
    code = (payload.item_code or "").strip()
    if not code:
        raise HTTPException(400, "กรุณาระบุรหัสอะไหล่")
    dup = s.exec(select(Item).where(func.lower(func.trim(Item.item_code)) == code.lower())).first()
    if dup:
        raise HTTPException(400, f"รหัสอะไหล่ {dup.item_code} มีอยู่แล้วในระบบ")
    data = payload.model_dump()
    data["item_code"] = code
    for k in ("category", "part_name", "part_number", "brand", "description",
              "machine_group", "uom", "box", "level"):
        data[k] = (data[k] or "").strip()
    if not data["description"]:
        data["description"] = " \\ ".join(x for x in (data["part_name"], data["part_number"], data["brand"]) if x)
    it = Item(**data, source_sheet="manual", updated_at=datetime.utcnow())
    s.add(it); s.flush()
    # baseline ledger row: lets "reset test transactions" restore this starting stock
    s.add(StockTxn(item_id=it.id, item_code=code, txn_type="receive", qty=it.on_hand or 0,
                   balance_after=it.on_hand or 0, ref="CREATE", note="สร้างอะไหล่ใหม่",
                   user_id=u.id))
    s.commit()
    audit_log("item", f"created new item {code}", user=u)
    return {"ok": True, "item": item_dict(it)}


# --------------------------- returns (leader/admin) -----------------------
class ReturnIn(BaseModel):
    line_id: int
    qty: float
    reason: str


@app.post("/api/returns")
async def create_return(payload: ReturnIn, s: Session = Depends(get_session),
                        u: User = Depends(require("leader"))):
    """Put withdrawn parts back into stock (e.g. wrong part taken) with a reason.
    The requisition keeps a remark so the history stays auditable."""
    ln = s.get(ReqLine, payload.line_id)
    if not ln:
        raise HTTPException(404, "ไม่พบรายการเบิก")
    reason = (payload.reason or "").strip()
    if not reason:
        raise HTTPException(400, "กรุณาระบุเหตุผลการคืน")
    if payload.qty <= 0:
        raise HTTPException(400, "จำนวนคืนต้องมากกว่า 0")
    req = s.get(Requisition, ln.requisition_id)
    if not req or req.status not in ("confirmed", "reconciled"):
        raise HTTPException(400, "คืนได้เฉพาะใบเบิกที่ยืนยันแล้ว")
    it = s.get(Item, ln.item_id)
    if not it:
        raise HTTPException(404, "ไม่พบอะไหล่")
    q = round(payload.qty, 6)
    # atomic: only succeeds if the total returned stays <= withdrawn qty
    res = s.execute(sa_update(ReqLine).where(
        ReqLine.id == ln.id,
        func.coalesce(ReqLine.returned_qty, 0) + q <= ReqLine.qty + 1e-9
    ).values(returned_qty=func.coalesce(ReqLine.returned_qty, 0) + q)
     .execution_options(synchronize_session=False))
    if res.rowcount != 1:
        s.rollback(); s.refresh(ln)
        remaining = round((ln.qty or 0) - (ln.returned_qty or 0), 6)
        raise HTTPException(400, f"คืนได้สูงสุด {remaining:g} (เบิก {ln.qty:g}, คืนแล้ว {ln.returned_qty or 0:g})")
    bal = change_stock(s, it.id, q)                        # atomic
    s.refresh(ln)
    who = u.full_name or u.username
    s.add(StockTxn(item_id=it.id, item_code=it.item_code, txn_type="return", qty=q,
                   balance_after=bal, ref=req.ref_no if req else "", user_id=u.id,
                   machine_name=req.machine_name if req else "", note=f"คืน: {reason}"))
    s.add(ReturnLog(requisition_id=ln.requisition_id, reqline_id=ln.id, item_id=it.id,
                    item_code=it.item_code, qty=q, reason=reason, returned_by=who))
    s.commit(); s.refresh(it)
    audit_log("return", f"{req.ref_no if req else ''} {it.item_code} +{payload.qty:g}: {reason}", user=u)
    await notify("stock.received",
                 f"↩️ คืนอะไหล่ {it.item_code} x{payload.qty:g} (ใบเบิก {req.ref_no if req else '-'})\n"
                 f"เหตุผล: {reason}\nโดย: {who} · คงเหลือ {it.on_hand:g}",
                 {"item": item_dict(it)})
    return {"ok": True, "on_hand": it.on_hand, "returned_qty": ln.returned_qty}


@app.get("/api/returns")
def list_returns(limit: int = 200, s: Session = Depends(get_session),
                 u: User = Depends(require("leader"))):
    rows = s.exec(select(ReturnLog).order_by(ReturnLog.created_at.desc()).limit(limit)).all()
    out = []
    for x in rows:
        req = s.get(Requisition, x.requisition_id)
        out.append({"id": x.id, "ref_no": req.ref_no if req else "", "item_code": x.item_code,
                    "qty": x.qty, "reason": x.reason, "by": x.returned_by, "at": x.created_at,
                    "requester": req.requester_name if req else "",
                    "machine": req.machine_name if req else ""})
    return out


# --------------------------- delete a requisition (admin) -----------------
@app.delete("/api/requisitions/{req_id}")
def delete_requisition(req_id: int, s: Session = Depends(get_session),
                       u: User = Depends(require("admin"))):
    """Remove a (test) requisition and give back the stock it took (net of any
    returns already made), so on-hand goes back to where it was."""
    req = s.get(Requisition, req_id)
    if not req:
        raise HTTPException(404, "ไม่พบใบเบิก")
    was = req.status
    if not claim(s, Requisition, req_id, was, "deleting"):   # only one delete may proceed
        raise HTTPException(409, "ใบเบิกนี้กำลังถูกลบ/ถูกลบไปแล้ว")
    lines = s.exec(select(ReqLine).where(ReqLine.requisition_id == req.id)).all()
    restored = []
    if was in ("confirmed", "reconciled"):
        for ln in lines:
            give_back = round((ln.qty or 0) - (ln.returned_qty or 0), 6)
            it = s.get(Item, ln.item_id)
            if it and give_back:
                change_stock(s, it.id, give_back)          # atomic
                restored.append(f"{it.item_code}+{give_back:g}")
    s.query(StockTxn).filter(StockTxn.ref == req.ref_no).delete(synchronize_session=False)
    s.query(ReturnLog).filter(ReturnLog.requisition_id == req.id).delete(synchronize_session=False)
    s.query(ReqLine).filter(ReqLine.requisition_id == req.id).delete(synchronize_session=False)
    images.drop_blobs(s, "req", req.id)
    ref = req.ref_no
    s.delete(req); s.commit()
    audit_log("delete", f"deleted requisition {ref}; stock restored: {', '.join(restored) or '-'}", user=u)
    return {"ok": True, "ref_no": ref, "restored": restored}


# --------------------------- shutdown -------------------------------------
@app.post("/api/shutdown")
def shutdown(u: User = Depends(require("admin"))):
    """Cleanly stop the program from the UI (admin only)."""
    import os, threading, time
    audit_log("shutdown", "program stopped from UI", user=u)

    def _stop():
        time.sleep(0.8)
        os._exit(0)
    threading.Thread(target=_stop, daemon=True).start()
    return {"ok": True}


@app.middleware("http")
async def always_fresh_frontend(request, call_next):
    """Browsers (especially phones) must re-check page files on every load, so an
    update is never hidden behind an old cached app.js / styles.css."""
    resp = await call_next(request)
    p = request.url.path
    if p == "/" or (p.endswith((".html", ".js", ".css")) and not p.startswith("/vendor/")):
        resp.headers["Cache-Control"] = "no-cache, must-revalidate"
    return resp


from .admin import router as admin_router
app.include_router(admin_router)
app.include_router(images.router)


# --------------------------- static frontend ------------------------------
app.mount("/", StaticFiles(directory=str(WEB_DIR), html=True), name="web")
