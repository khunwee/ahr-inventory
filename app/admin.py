"""Administration tools: system info, backup/restore, data-quality checks,
clean-up actions and guarded data resets. All endpoints require the admin role.
Works on both SQLite (local install) and PostgreSQL (Neon / cloud)."""
import io
import os
import re
import shutil
import sys
import tempfile
import time
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from fastapi.responses import StreamingResponse, FileResponse
from pydantic import BaseModel
from sqlalchemy import select, func, text, LargeBinary
import base64

from .auth import require, verify_pw
from .audit import log as audit_log
from .config import settings, BASE_DIR, DATA_DIR
from .db import engine, SessionLocal, get_session, Session, _is_sqlite, init_db
from .models import (Base, User, Machine, Item, Requisition, ReqLine, StockTxn,
                     PurchaseOrder, POLine, ReturnLog, AuditLog, Setting, ImageBlob)
from .tz import local_now, fmt_local
from .services import change_stock

router = APIRouter(prefix="/api/admin", tags=["admin"])
ADMIN = require("admin")
STARTED = time.time()
CONFIRM_PHRASE = "ยืนยันล้างข้อมูล"
BASELINE_REFS = ("OPENING", "IMPORT", "CREATE")   # ledger rows that define starting stock
INVENTORY_TABLES = (ImageBlob, ReturnLog, POLine, PurchaseOrder, ReqLine, Requisition, StockTxn, Item, Machine)
CHUNK = 30000          # Excel cell limit is 32,767 chars -> binary data is split over rows


def _counts(s):
    out = {}
    for m in (Item, Machine, Requisition, ReqLine, ReturnLog, StockTxn, PurchaseOrder, User, AuditLog, ImageBlob):
        out[m.__tablename__] = s.execute(select(func.count()).select_from(m)).scalar_one()
    return out


def _db_size_bytes(s):
    try:
        if _is_sqlite:
            p = Path(settings.DB_URL.replace("sqlite:///", ""))
            return p.stat().st_size if p.exists() else 0
        return s.execute(text("select pg_database_size(current_database())")).scalar_one()
    except Exception:
        return None


# ------------------------------------------------------------------ system --
@router.get("/system")
def system_info(s: Session = Depends(get_session), u: User = Depends(ADMIN)):
    from .optimize import solver_status
    weak = []
    for x in s.exec(select(User).where(User.active == True)).all():
        for pw in ("admin123", "user123"):
            if verify_pw(pw, x.password_hash):
                weak.append(x.username)
    try:
        du = shutil.disk_usage(str(BASE_DIR))
        disk = {"free_gb": round(du.free / 1e9, 1), "total_gb": round(du.total / 1e9, 1)}
    except Exception:
        disk = None
    backups = []
    bdir = BASE_DIR / "backups"
    if bdir.exists():
        backups = [{"name": f.name, "size": f.stat().st_size,
                    "at": fmt_local(datetime.fromtimestamp(f.stat().st_mtime, timezone.utc).replace(tzinfo=None))}
                   for f in sorted(bdir.glob("inventory_*.*"), reverse=True)[:10]]
    return {
        "db_type": "SQLite (ไฟล์ในเครื่อง)" if _is_sqlite else "PostgreSQL (ฐานข้อมูลออนไลน์)",
        "db_size": _db_size_bytes(s), "counts": _counts(s),
        "server_time": local_now().strftime("%d/%m/%Y %H:%M:%S"), "timezone": settings.TIMEZONE,
        "uptime_min": round((time.time() - STARTED) / 60, 1),
        "python": sys.version.split()[0], "solvers": solver_status(),
        "disk": disk, "weak_password_users": weak, "server_backups": backups,
        "image_bytes": s.execute(select(func.coalesce(func.sum(ImageBlob.size), 0))).scalar_one(),
        "items_with_photo": s.execute(select(func.count()).select_from(Item).where(Item.image_ver > 0)).scalar_one(),
        "is_sqlite": _is_sqlite,
    }


# ------------------------------------------------------------------ backup --
def _workbook_bytes():
    """Every table -> one sheet. Restorable with /restore."""
    from openpyxl import Workbook
    wb = Workbook()
    meta = wb.active; meta.title = "_meta"
    meta.append(["app", settings.APP_NAME]); meta.append(["created", local_now().strftime("%d/%m/%Y %H:%M:%S")])
    meta.append(["format", "ahr-backup-v1"])
    with engine.connect() as conn:
        for table in Base.metadata.sorted_tables:
            ws = wb.create_sheet(table.name[:31])
            cols = [c.name for c in table.columns]
            bins = [c.name for c in table.columns if isinstance(c.type, LargeBinary)]
            ws.append(cols + (["__part"] if bins else []))
            for row in conn.execute(select(table)).mappings():
                if not bins:
                    ws.append([row[c] for c in cols]); continue
                enc = base64.b64encode(row[bins[0]] or b"").decode()
                parts = [enc[i:i + CHUNK] for i in range(0, len(enc), CHUNK)] or [""]
                for k, part in enumerate(parts):
                    vals = [row[c] if (k == 0 and c not in bins) else (part if c == bins[0] else (row[c] if c == "id" else None))
                            for c in cols]
                    ws.append(vals + [k])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


@router.get("/backup.xlsx")
def backup_xlsx(u: User = Depends(ADMIN)):
    buf = _workbook_bytes()
    audit_log("backup", "downloaded full backup (xlsx)", user=u)
    fname = f"AHR_backup_{local_now():%Y%m%d_%H%M}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={fname}"})


@router.get("/backup.db")
def backup_db(u: User = Depends(ADMIN)):
    if not _is_sqlite:
        raise HTTPException(400, "ฐานข้อมูลออนไลน์ (PostgreSQL) ให้ใช้ไฟล์สำรองแบบ Excel แทน")
    import sqlite3
    src = settings.DB_URL.replace("sqlite:///", "")
    tmp = Path(tempfile.gettempdir()) / f"ahr_backup_{int(time.time())}.db"
    with sqlite3.connect(src) as a, sqlite3.connect(tmp) as b:
        a.backup(b)                      # consistent copy even while the app is running
    audit_log("backup", "downloaded database file (.db)", user=u)
    return FileResponse(tmp, filename=f"AHR_inventory_{local_now():%Y%m%d_%H%M}.db",
                        media_type="application/octet-stream")


def _snapshot_to_disk(tag):
    """Safety copy written before destructive operations."""
    try:
        d = BASE_DIR / "backups"; d.mkdir(exist_ok=True)
        p = d / f"inventory_{tag}_{local_now():%Y%m%d_%H%M%S}.xlsx"
        p.write_bytes(_workbook_bytes().getvalue())
        for old in sorted(d.glob("inventory_pre-*.xlsx"), key=lambda f: f.stat().st_mtime)[:-10]:
            try:
                old.unlink()                      # keep only the 10 newest safety copies
            except Exception:
                pass
        return p.name
    except Exception:
        return None


def _reset_sequences(conn):
    if _is_sqlite:
        return
    prep = engine.dialect.identifier_preparer
    for table in Base.metadata.sorted_tables:
        pk = [c for c in table.primary_key.columns]
        if len(pk) == 1 and str(pk[0].type).upper().startswith("INTEGER"):
            t = prep.quote(table.name)
            conn.execute(text(
                f"SELECT setval(pg_get_serial_sequence('{t}', '{pk[0].name}'), "
                f"COALESCE((SELECT MAX({pk[0].name}) FROM {t}), 0) + 1, false)"))


@router.post("/restore")
async def restore(file: UploadFile = File(...), confirm_text: str = Form(""),
                  u: User = Depends(ADMIN)):
    if confirm_text.strip() != CONFIRM_PHRASE:
        raise HTTPException(400, f"พิมพ์ '{CONFIRM_PHRASE}' เพื่อยืนยันการกู้คืน")
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(await file.read()), data_only=True)
    except Exception:
        raise HTTPException(400, "ไฟล์ไม่ถูกต้อง — ต้องเป็นไฟล์สำรอง .xlsx ที่ดาวน์โหลดจากระบบนี้")
    if "_meta" not in wb.sheetnames or "item" not in wb.sheetnames:
        raise HTTPException(400, "ไม่ใช่ไฟล์สำรองของระบบนี้ (ไม่พบชีต _meta / item)")
    safety = _snapshot_to_disk("pre-restore")
    tables = [t for t in Base.metadata.sorted_tables if t.name in wb.sheetnames]
    restored = {}
    with engine.begin() as conn:
        for t in reversed(Base.metadata.sorted_tables):       # children first
            if t.name in wb.sheetnames:
                conn.execute(t.delete())
        for t in tables:                                       # parents first
            ws = wb[t.name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            header = [h for h in rows[0]]
            valid = {c.name for c in t.columns}
            bins = [c.name for c in t.columns if isinstance(c.type, LargeBinary)]
            if bins and "__part" in header:                      # join split binary rows
                hi = {h: i for i, h in enumerate(header)}
                joined = {}
                for r in rows[1:]:
                    rid = r[hi["id"]]
                    if rid is None:
                        continue
                    if (r[hi["__part"]] or 0) == 0:
                        joined[rid] = list(r)
                    elif rid in joined:
                        joined[rid][hi[bins[0]]] = (joined[rid][hi[bins[0]]] or "") + (r[hi[bins[0]]] or "")
                for rec in joined.values():
                    rec[hi[bins[0]]] = base64.b64decode(rec[hi[bins[0]]] or "")
                rows = [rows[0]] + [tuple(v) for v in joined.values()]
            batch = []
            for r in rows[1:]:
                rec = {h: v for h, v in zip(header, r) if h in valid}
                for c in t.columns:            # coerce to the column's python type
                    v = rec.get(c.name)
                    if v is None:
                        continue
                    try:
                        pt = c.type.python_type
                    except Exception:
                        continue
                    if pt is bytes:
                        continue
                    if pt is str and not isinstance(v, str):
                        rec[c.name] = str(v)
                    elif pt is bool:
                        rec[c.name] = bool(v) if not isinstance(v, str) else v.lower() in ("1", "true")
                    elif pt is float and isinstance(v, (int, str)):
                        rec[c.name] = float(v)
                    elif pt is int and isinstance(v, (float, str)):
                        rec[c.name] = int(float(v))
                batch.append(rec)
            if batch:
                conn.execute(t.insert(), batch)
            restored[t.name] = len(batch)
        _reset_sequences(conn)
    audit_log("restore", f"restored from backup {file.filename}: {restored}", user=u)
    return {"ok": True, "restored": restored, "safety_copy": safety}


# ------------------------------------------------ ledger (stock history) ----
def _signed(t):
    """Effect of a non-baseline ledger row on on-hand."""
    q = t.qty or 0
    return -q if t.txn_type == "issue" else q          # receive / return / adjust(+/-)


def _ledger_state(s):
    """Per item: value expected from the ledger = latest baseline (opening /
    import / create) + every movement after it."""
    st = {}
    for t in s.exec(select(StockTxn).order_by(StockTxn.created_at, StockTxn.id)).all():
        e = st.setdefault(t.item_id, {"base": 0.0, "after": 0.0, "openings": [], "later_base": False})
        if t.ref in BASELINE_REFS:
            if t.ref == "OPENING":
                e["openings"].append(t)
            else:
                e["later_base"] = True
            e["base"] = t.balance_after or 0.0
            e["after"] = 0.0
        else:
            e["after"] += _signed(t)
    return st


def _expected(e):
    return round((e["base"] if e else 0.0) + (e["after"] if e else 0.0), 6)


# ------------------------------------------------------------------ health --
def _norm(v):
    return re.sub(r"\s+", " ", (v or "").strip()).lower()


def _uom_key(v):
    return (v or "").strip().strip(".").lower()


@router.get("/health")
def health(s: Session = Depends(get_session), u: User = Depends(ADMIN)):
    items = s.exec(select(Item)).all()
    machines = s.exec(select(Machine)).all()
    checks = []

    def add(key, title, severity, rows, fix=None, hint=""):
        checks.append({"key": key, "title": title, "severity": severity, "count": len(rows),
                       "samples": rows[:25], "fix": fix if rows else None, "hint": hint})

    # 1) item codes that differ only by case / spaces
    g = defaultdict(list)
    for it in items:
        g[_norm(it.item_code)].append(it)
    rows = [{"text": " , ".join(x.item_code for x in v), "ids": [x.id for x in v]} for v in g.values() if len(v) > 1]
    add("dup_code", "รหัสอะไหล่ซ้ำ (ต่างกันแค่ตัวพิมพ์/ช่องว่าง)", "error", rows,
        hint="ตรวจสอบแล้วลบ/แก้ไขรายการที่ซ้ำในหน้าคลังอะไหล่")
    # 2) same part number under different codes
    g = defaultdict(list)
    for it in items:
        if _norm(it.part_number) and len(_norm(it.part_number)) >= 3:
            g[(_norm(it.part_number), _norm(it.brand))].append(it)
    rows = [{"text": f"เบอร์ {v[0].part_number} ({v[0].brand or '-'}): " + " , ".join(x.item_code for x in v),
             "ids": [x.id for x in v]} for v in g.values() if len(v) > 1]
    add("dup_partno", "อาจเป็นอะไหล่ตัวเดียวกัน (เบอร์อะไหล่+ยี่ห้อ ซ้ำ แต่คนละรหัส)", "warn", rows,
        hint="ถ้าเป็นชิ้นเดียวกันจริง ควรรวมเป็นรหัสเดียว")
    # 3) same description under different codes
    g = defaultdict(list)
    for it in items:
        if _norm(it.description):
            g[_norm(it.description)].append(it)
    rows = [{"text": f"{v[0].description[:50]}: " + " , ".join(x.item_code for x in v),
             "ids": [x.id for x in v]} for v in g.values() if len(v) > 1]
    add("dup_desc", "รายละเอียดเหมือนกันทุกตัวอักษร แต่คนละรหัส", "warn", rows)
    # 4) whitespace problems
    fields = ("item_code", "part_name", "part_number", "brand", "machine_group", "uom", "box", "level", "category")
    rows = []
    for it in items:
        bad = [f for f in fields if (getattr(it, f) or "") != re.sub(r"\s+", " ", (getattr(it, f) or "").strip())]
        if bad:
            rows.append({"text": f"{it.item_code}: {', '.join(bad)}", "ids": [it.id]})
    add("whitespace", "ข้อความมีช่องว่างเกิน (หน้า/ท้าย/ซ้อน)", "warn", rows, fix="trim_whitespace")
    # 5) UOM spelled several ways
    g = defaultdict(set)
    for it in items:
        if it.uom:
            g[_uom_key(it.uom)].add(it.uom)
    rows = [{"text": " / ".join(sorted(v))} for v in g.values() if len(v) > 1]
    add("uom_variants", "หน่วยนับเขียนหลายแบบ (เช่น Pcs / PCS / pcs.)", "warn", rows, fix="normalize_uom")
    # 6) negative stock
    rows = [{"text": f"{it.item_code}: {it.on_hand:g}", "ids": [it.id]} for it in items if (it.on_hand or 0) < 0]
    add("negative", "ยอดคงเหลือติดลบ (เบิกเกินสต็อก)", "error", rows, fix="fix_negative_stock",
        hint="ควรตรวจนับของจริงก่อน — ปุ่มแก้ไขจะตั้งยอดเป็น 0 และบันทึกลง ledger")
    # 7) missing data
    add("no_price", "ไม่มีราคาต่อหน่วย (ราคา = 0)", "info",
        [{"text": it.item_code, "ids": [it.id]} for it in items if not (it.unit_price or 0)],
        hint="มูลค่าคงคลังและ Optimization จะคำนวณไม่ครบ")
    add("no_name", "ไม่มีชื่ออะไหล่", "warn",
        [{"text": it.item_code, "ids": [it.id]} for it in items if not (it.part_name or "").strip()])
    add("no_location", "ไม่ระบุที่เก็บ (box)", "info",
        [{"text": it.item_code, "ids": [it.id]} for it in items if not (it.box or "").strip()])
    add("no_uom", "ไม่ระบุหน่วยนับ", "warn",
        [{"text": it.item_code, "ids": [it.id]} for it in items if not (it.uom or "").strip()], fix="fill_uom")
    # 8) min / ROP / max inconsistent
    rows = [{"text": f"{it.item_code}: ROP {it.reorder_point:g} > Max {it.max_level:g}", "ids": [it.id]}
            for it in items if (it.max_level or 0) > 0 and (it.reorder_point or 0) > (it.max_level or 0)]
    add("rop_gt_max", "จุดสั่งซื้อ (ROP) สูงกว่าระดับสูงสุด (Max)", "warn", rows,
        hint="ใช้โหมด 'แนะนำค่า Min/Max' ในหน้า Optimize เพื่อปรับให้เหมาะสม")
    # 9) duplicate machine names
    g = defaultdict(list)
    for m in machines:
        g[_norm(m.name)].append(m)
    rows = [{"text": " , ".join(x.name for x in v)} for v in g.values() if len(v) > 1]
    add("dup_machine", "ชื่อเครื่องจักรซ้ำ (ต่างกันแค่ตัวพิมพ์/ช่องว่าง)", "warn", rows, fix="merge_machines")
    # 10) orphans
    item_ids = {it.id for it in items}
    orphan = [x for x in s.exec(select(ReqLine)).all() if x.item_id not in item_ids]
    add("orphan_lines", "รายการเบิกที่อ้างถึงอะไหล่ที่ไม่มีแล้ว", "info",
        [{"text": f"line {x.id}: {x.item_code}"} for x in orphan])
    # 11) on-hand vs stock-movement history (catches any calculation error)
    st = _ledger_state(s)
    rows = []
    for it in items:
        exp = _expected(st.get(it.id))
        if abs((it.on_hand or 0) - exp) > 1e-6:
            rows.append({"text": f"{it.item_code}: ระบบ {it.on_hand:g} / ตามประวัติ {exp:g}", "ids": [it.id]})
    add("ledger_mismatch", "ยอดคงเหลือไม่ตรงกับประวัติความเคลื่อนไหว (เบิก/รับ/คืน/ปรับยอด)", "error", rows,
        fix="recalc_from_ledger", hint="แก้ไขอัตโนมัติ = คำนวณยอดคงเหลือใหม่จากประวัติทุกรายการ")
    # 12) parts listed on several rows of the source Excel
    old_dup, merged_ok = [], []
    for it in items:
        e = st.get(it.id)
        if not e:
            continue
        ops = e["openings"]
        if len(ops) > 1 and not e["later_base"]:
            old_dup.append({"text": f"{it.item_code}: " + " + ".join(f"{o.qty:g}" for o in ops) + f" (ระบบ {it.on_hand:g})",
                            "ids": [it.id]})
        elif ops and (ops[-1].note or "").startswith("รวม"):
            merged_ok.append({"text": f"{it.item_code} = {ops[-1].qty:g} [{it.box}]", "ids": [it.id]})
    add("opening_dup", "นำเข้าด้วยเวอร์ชันเก่า: อะไหล่ที่มีหลายแถวในไฟล์ต้นฉบับถูกนับยอดไม่ครบ", "error", old_dup,
        fix="fix_opening_duplicates", hint="แก้ไขอัตโนมัติ = รวมยอดทุกแถวในไฟล์ต้นฉบับ (คงผลการเบิก/รับ/คืนที่เกิดขึ้นหลังจากนั้นไว้)")
    add("merged_rows", "อะไหล่ที่รวมยอดจากหลายแถวในไฟล์ต้นฉบับ (ควรตรวจนับของจริงยืนยัน)", "info", merged_ok,
        hint="ไฟล์ Excel เดิมระบุรหัสเดียวกันหลายแถว (เช่น เก็บ 2 ที่) ระบบรวมยอดให้แล้ว")
    # 13) default passwords
    weak = []
    for x in s.exec(select(User).where(User.active == True)).all():
        if any(verify_pw(pw, x.password_hash) for pw in ("admin123", "user123")):
            weak.append({"text": f"{x.username} ({x.role})"})
    add("weak_pw", "บัญชีที่ยังใช้รหัสผ่านเริ่มต้น (ความปลอดภัย)", "error", weak,
        hint="เมนูผู้ใช้งาน → แก้ไข → ตั้งรหัสใหม่ หรือปิดใช้งานบัญชีทดลอง")
    summary = {sev: sum(1 for c in checks if c["severity"] == sev and c["count"]) for sev in ("error", "warn", "info")}
    return {"checked_items": len(items), "summary": summary, "checks": checks}


class CleanupIn(BaseModel):
    action: str


@router.post("/cleanup")
def cleanup(payload: CleanupIn, s: Session = Depends(get_session), u: User = Depends(ADMIN)):
    n = 0
    items = s.exec(select(Item)).all()
    if payload.action == "trim_whitespace":
        fields = ("item_code", "part_name", "part_number", "brand", "machine_group", "uom", "box", "level", "category")
        codes = {it.item_code for it in items}
        for it in items:
            changed = False
            for f in fields:
                v = getattr(it, f) or ""
                nv = re.sub(r"\s+", " ", v.strip())
                if f == "item_code" and nv != v and nv in codes:
                    continue                      # would collide — leave for manual review
                if nv != v:
                    setattr(it, f, nv); changed = True
            if changed:
                n += 1
    elif payload.action == "normalize_uom":
        from scripts.import_excel import clean_uom
        for it in items:
            nv = clean_uom(it.uom)
            if nv != it.uom:
                it.uom = nv; n += 1
    elif payload.action == "fill_uom":
        for it in items:
            if not (it.uom or "").strip():
                it.uom = "Pcs"; n += 1
    elif payload.action == "fix_negative_stock":
        for it in items:
            if (it.on_hand or 0) < 0:
                d = -(it.on_hand or 0)
                bal = change_stock(s, it.id, d)
                s.add(StockTxn(item_id=it.id, item_code=it.item_code, txn_type="adjust",
                               qty=d, balance_after=bal, ref="CLEANUP",
                               note="ปรับยอดติดลบเป็น 0", user_id=u.id)); n += 1
    elif payload.action == "merge_machines":
        g = defaultdict(list)
        for m in s.exec(select(Machine).order_by(Machine.id)).all():
            g[_norm(m.name)].append(m)
        for group in g.values():
            keep, dups = group[0], group[1:]
            for d in dups:
                s.query(Requisition).filter(Requisition.machine_id == d.id).update(
                    {"machine_id": keep.id}, synchronize_session=False)
                s.delete(d); n += 1
    elif payload.action == "fix_opening_duplicates":
        st = _ledger_state(s)
        for it in items:
            e = st.get(it.id)
            if not e or len(e["openings"]) < 2 or e["later_base"]:
                continue
            ops = e["openings"]
            total = round(sum(o.qty or 0 for o in ops), 6)          # true opening = sum of rows
            imported = round((it.on_hand or 0) - e["after"], 6)      # what the old import set
            d = round(total - imported, 6)
            first_at = ops[0].created_at
            note = "รวม %d แถวจากไฟล์ต้นฉบับ (แก้ไขยอด): %s" % (len(ops), " + ".join(f"{o.qty:g}" for o in ops))
            for o in ops:
                s.delete(o)
            s.flush()
            s.add(StockTxn(item_id=it.id, item_code=it.item_code, txn_type="receive", qty=total,
                           balance_after=total, ref="OPENING", note=note, created_at=first_at))
            if d:
                change_stock(s, it.id, d)
            n += 1
    elif payload.action == "recalc_from_ledger":
        st = _ledger_state(s)
        for it in items:
            exp = _expected(st.get(it.id))
            d = round(exp - (it.on_hand or 0), 6)
            if abs(d) > 1e-6:
                change_stock(s, it.id, d); n += 1
    else:
        raise HTTPException(400, "unknown action")
    s.commit()
    audit_log("cleanup", f"{payload.action}: {n} record(s)", user=u)
    return {"ok": True, "action": payload.action, "fixed": n}


@router.delete("/items/{item_id}")
def delete_item(item_id: int, s: Session = Depends(get_session), u: User = Depends(ADMIN)):
    it = s.get(Item, item_id)
    if not it:
        raise HTTPException(404, "ไม่พบอะไหล่")
    used = s.execute(select(func.count()).select_from(ReqLine).where(ReqLine.item_id == item_id)).scalar_one() \
        + s.execute(select(func.count()).select_from(POLine).where(POLine.item_id == item_id)).scalar_one()
    if used:
        raise HTTPException(400, "อะไหล่นี้มีประวัติการเบิก/สั่งซื้อ จึงลบไม่ได้ — แก้ไขข้อมูลแทน "
                                 "(หรือลบใบเบิกทดสอบที่เกี่ยวข้องก่อน)")
    code = it.item_code
    s.query(StockTxn).filter(StockTxn.item_id == item_id).delete(synchronize_session=False)
    s.query(ImageBlob).filter(ImageBlob.owner_type == "item", ImageBlob.owner_id == item_id).delete(synchronize_session=False)
    s.delete(it); s.commit()
    audit_log("item", f"deleted item {code}", user=u)
    return {"ok": True}


# ------------------------------------------------------------------ reset ---
class ResetIn(BaseModel):
    mode: str             # transactions | factory | empty
    confirm_text: str


@router.post("/reset")
def reset(payload: ResetIn, u: User = Depends(ADMIN)):
    if payload.confirm_text.strip() != CONFIRM_PHRASE:
        raise HTTPException(400, f"พิมพ์ '{CONFIRM_PHRASE}' ให้ถูกต้องเพื่อยืนยัน")
    if payload.mode not in ("transactions", "factory", "empty"):
        raise HTTPException(400, "unknown mode")
    safety = _snapshot_to_disk(f"pre-{payload.mode}")
    result = {}
    with SessionLocal() as s:
        if payload.mode == "transactions":
            # Undo precisely the stock movements we delete: for each item, every
            # non-baseline ledger row AFTER its latest baseline (import/create) is
            # reversed; rows before that baseline were already superseded by it.
            SIGN = {"issue": -1, "receive": 1, "return": 1, "adjust": 1}
            last_base = {}
            txns = s.exec(select(StockTxn).order_by(StockTxn.created_at, StockTxn.id)).all()
            for t in txns:
                if t.ref in BASELINE_REFS:
                    last_base[t.item_id] = (t.created_at, t.id)
            delta = defaultdict(float)
            for t in txns:
                if t.ref in BASELINE_REFS:
                    continue
                b = last_base.get(t.item_id)
                if b is None or (t.created_at, t.id) > b:
                    delta[t.item_id] += SIGN.get(t.txn_type, 0) * (t.qty or 0)
            counts = {"requisitions": s.execute(select(func.count()).select_from(Requisition)).scalar_one(),
                      "purchase_orders": s.execute(select(func.count()).select_from(PurchaseOrder)).scalar_one(),
                      "returns": s.execute(select(func.count()).select_from(ReturnLog)).scalar_one()}
            for m in (ReturnLog, POLine, PurchaseOrder, ReqLine, Requisition):
                s.query(m).delete(synchronize_session=False)
            s.query(ImageBlob).filter(ImageBlob.owner_type == "req").delete(synchronize_session=False)
            s.query(StockTxn).filter(~StockTxn.ref.in_(BASELINE_REFS)).delete(synchronize_session=False)
            restored = 0
            for it in s.exec(select(Item)).all():
                d = round(delta.get(it.id, 0.0), 6)
                if d:
                    change_stock(s, it.id, -d); restored += 1
            s.commit()
            result = {**counts, "items_stock_restored": restored}
        else:
            for m in INVENTORY_TABLES:
                s.query(m).delete(synchronize_session=False)
            s.commit()
            result = {"cleared": True}
    if payload.mode == "factory":
        from scripts.bootstrap import run as bootstrap_run
        bootstrap_run()
        with SessionLocal() as s:
            result["items_loaded"] = s.execute(select(func.count()).select_from(Item)).scalar_one()
    if not _is_sqlite:
        with engine.begin() as conn:
            _reset_sequences(conn)
    audit_log("reset", f"mode={payload.mode} result={result}", user=u)
    return {"ok": True, "mode": payload.mode, "result": result, "safety_copy": safety}


@router.post("/clear-audit")
def clear_audit(payload: ResetIn, u: User = Depends(ADMIN)):
    if payload.confirm_text.strip() != CONFIRM_PHRASE:
        raise HTTPException(400, f"พิมพ์ '{CONFIRM_PHRASE}' ให้ถูกต้องเพื่อยืนยัน")
    with SessionLocal() as s:
        n = s.query(AuditLog).delete(synchronize_session=False); s.commit()
    audit_log("reset", f"audit log cleared ({n} rows)", user=u)
    return {"ok": True, "deleted": n}
