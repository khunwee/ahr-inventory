"""Template import with a duplicate check BEFORE anything is written.

Flow: the admin uploads the file to /api/import/preview -> gets a report
(new parts, parts already in the system, changed values, duplicate rows in
the file, possible duplicates under a different code, invalid rows) -> picks
how to import -> /api/import/upload applies exactly that choice.

Modes
  new_only : add parts whose code is not in the system; existing parts untouched
  update   : add new parts + update the details of existing parts that changed
             (on-hand is NOT changed unless update_stock=True, so re-uploading an
             old file can never wipe out withdrawals made since)
  replace  : clear inventory data, then load the file as a fresh dataset
Duplicate rows of one code inside the file
  first    : use the first row, skip the others (default)
  sum      : one part, on-hand = sum of the rows (part kept in several places)
"""
import io
import re
from datetime import datetime

from sqlalchemy import select, func

from .models import (Item, Machine, StockTxn, Requisition, ReqLine, ReturnLog,
                     PurchaseOrder, POLine, ImageBlob)
from .services import change_stock

IMPORT_COLS = ["item_code", "category", "description", "part_name", "part_number",
               "brand", "machine_group", "uom", "box", "level", "unit_price",
               "on_hand", "min_level", "reorder_point", "lead_time_months"]
TEXT_FIELDS = ["category", "description", "part_name", "part_number", "brand",
               "machine_group", "uom", "box", "level"]
NUM_FIELDS = ["unit_price", "on_hand", "min_level", "reorder_point", "lead_time_months"]
DETAIL_FIELDS = TEXT_FIELDS + ["unit_price", "min_level", "reorder_point", "lead_time_months"]
LABEL = {"category": "หมวด", "description": "รายละเอียด", "part_name": "ชื่อ", "part_number": "เบอร์อะไหล่",
         "brand": "ยี่ห้อ", "machine_group": "เครื่อง", "uom": "หน่วย", "box": "กล่อง", "level": "ชั้น",
         "unit_price": "ราคา", "on_hand": "คงเหลือ", "min_level": "Min", "reorder_point": "ROP",
         "lead_time_months": "Lead"}
MAX_LIST = 500


class ImportError_(Exception):
    pass


def norm_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)                                   # Excel turns 1 into 1.0
    return re.sub(r"\s+", " ", str(v)).strip()


def code_key(v):
    return norm_text(v).upper()


def parse_num(v, default):
    """-> (value, error). Accepts 1,500 / ' 12 ' / empty."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return default, None
    if isinstance(v, bool):
        return default, "ไม่ใช่ตัวเลข"
    if isinstance(v, (int, float)):
        f = float(v)
    else:
        try:
            f = float(str(v).replace(",", "").strip())
        except ValueError:
            return default, f"'{v}' ไม่ใช่ตัวเลข"
    if f < 0:
        return default, f"ค่าติดลบ ({f:g})"
    return f, None


def parse_workbook(data: bytes):
    from openpyxl import load_workbook
    if len(data) > 20 * 1024 * 1024:
        raise ImportError_("ไฟล์ใหญ่เกิน 20 MB — แบ่งไฟล์เป็นหลายส่วนแล้วนำเข้าทีละไฟล์")
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        raise ImportError_("ไฟล์ไม่ถูกต้อง — ต้องเป็นไฟล์ Excel (.xlsx) ตาม Template")
    rows, invalid, machines = [], [], []
    if "Items" in wb.sheetnames:
        it = wb["Items"].iter_rows(values_only=True)
        header = [norm_text(c).lower() for c in (next(it, None) or [])]
        idx = {n: header.index(n) for n in IMPORT_COLS if n in header}
        if "item_code" not in idx:
            raise ImportError_("ไม่พบคอลัมน์ item_code ในชีต Items (แถวแรกต้องเป็นหัวคอลัมน์ตาม Template)")
        for rn, row in enumerate(it, start=2):
            row = list(row or [])
            cell = lambda k: row[idx[k]] if k in idx and idx[k] < len(row) else None
            if all(norm_text(c) == "" for c in row):
                continue                                    # blank line
            code = norm_text(cell("item_code"))
            rec = {"row": rn, "item_code": code}
            for k in TEXT_FIELDS:
                rec[k] = norm_text(cell(k))
            rec["uom"] = rec["uom"] or "Pcs"
            errs = []
            if not code:
                errs.append("ไม่มี item_code")
            for k in NUM_FIELDS:
                v, e = parse_num(cell(k), 2.0 if k == "lead_time_months" else 0.0)
                rec[k] = v
                if e:
                    errs.append(f"{k}: {e}")
            rec["_has"] = {k for k in IMPORT_COLS if k in idx}
            if errs:
                invalid.append({"row": rn, "item_code": code, "errors": errs})
            else:
                rows.append(rec)
    if "Machines" in wb.sheetnames:
        for rn, row in enumerate(wb["Machines"].iter_rows(min_row=2, values_only=True), start=2):
            name = norm_text(row[0]) if row else ""
            if name:
                machines.append({"row": rn, "name": name})
    if not rows and not invalid and not machines:
        raise ImportError_("ไม่พบข้อมูลในชีต Items หรือ Machines")
    return {"rows": rows, "invalid": invalid, "machines": machines}


def _same(a, b):
    if isinstance(a, float) or isinstance(b, float):
        try:
            return abs(float(a or 0) - float(b or 0)) < 1e-9
        except (TypeError, ValueError):
            return False
    return norm_text(a) == norm_text(b)


def _group(rows):
    g = {}
    for r in rows:
        g.setdefault(code_key(r["item_code"]), []).append(r)
    return g


def _merge_group(grp, dup_rows):
    first = dict(grp[0])
    if len(grp) > 1 and dup_rows == "sum":
        first["on_hand"] = round(sum(r["on_hand"] for r in grp), 6)
        locs = []
        for r in grp:
            loc = " ".join(x for x in (r["box"], r["level"]) if x)
            if loc and loc not in locs:
                locs.append(loc)
        if len(locs) > 1:
            first["box"], first["level"] = " + ".join(locs), ""
    return first


def analyze(s, parsed):
    rows = parsed["rows"]
    groups = _group(rows)
    existing = {code_key(i.item_code): i for i in s.execute(select(Item)).scalars()}

    in_file = []
    for k, grp in groups.items():
        if len(grp) > 1:
            identical = all(all(_same(g[f], grp[0][f]) for f in IMPORT_COLS if f != "item_code") for g in grp[1:])
            in_file.append({"item_code": grp[0]["item_code"], "rows": [g["row"] for g in grp],
                            "identical": identical,
                            "on_hand": [g["on_hand"] for g in grp]})

    new, same, changed = [], [], []
    for k, grp in groups.items():
        r = grp[0]
        it = existing.get(k)
        if not it:
            new.append({"row": r["row"], "item_code": r["item_code"], "name": r["part_name"] or r["description"][:40],
                        "on_hand": r["on_hand"]})
            continue
        diffs = []
        for f in DETAIL_FIELDS:
            if f in r["_has"] and not _same(getattr(it, f), r[f]):
                diffs.append({"field": f, "label": LABEL[f], "old": getattr(it, f), "new": r[f]})
        stock_diff = None
        if "on_hand" in r["_has"] and not _same(it.on_hand, r["on_hand"]):
            stock_diff = {"old": it.on_hand, "new": r["on_hand"]}
        rec = {"row": r["row"], "item_code": it.item_code, "name": it.part_name or (it.description or "")[:40],
               "changes": diffs, "stock": stock_diff,
               "code_case_differs": it.item_code != r["item_code"]}
        (changed if diffs or stock_diff else same).append(rec)

    # possible duplicates: a NEW code whose part number+brand / description already exists
    def sig_pn(x):
        pn = norm_text(x.get("part_number") if isinstance(x, dict) else x.part_number).upper()
        br = norm_text(x.get("brand") if isinstance(x, dict) else x.brand).upper()
        return (pn, br) if len(pn) >= 3 else None

    def sig_desc(x):
        d = norm_text(x.get("description") if isinstance(x, dict) else x.description).upper()
        return d if len(d) >= 8 else None

    by_pn, by_desc = {}, {}
    for it in existing.values():
        if sig_pn(it):
            by_pn.setdefault(sig_pn(it), []).append(it.item_code)
        if sig_desc(it):
            by_desc.setdefault(sig_desc(it), []).append(it.item_code)
    new_keys = {code_key(n["item_code"]) for n in new}
    file_pn, file_desc = {}, {}
    for k in new_keys:
        r = groups[k][0]
        if sig_pn(r):
            file_pn.setdefault(sig_pn(r), []).append(r["item_code"])
        if sig_desc(r):
            file_desc.setdefault(sig_desc(r), []).append(r["item_code"])
    possible = []
    for k in new_keys:
        r = groups[k][0]
        why, others = [], []
        p, d = sig_pn(r), sig_desc(r)
        if p and by_pn.get(p):
            why.append(f"เบอร์ {r['part_number']} ({r['brand'] or '-'}) มีในระบบแล้ว"); others += by_pn[p]
        elif d and by_desc.get(d):
            why.append("รายละเอียดเหมือนอะไหล่ในระบบ"); others += by_desc[d]
        if p and len(file_pn.get(p, [])) > 1:
            why.append("เบอร์อะไหล่ซ้ำกับแถวอื่นในไฟล์"); others += [c for c in file_pn[p] if c != r["item_code"]]
        if why:
            possible.append({"row": r["row"], "item_code": r["item_code"], "reason": "; ".join(why),
                             "similar": sorted(set(others))[:6]})

    exist_m = {norm_text(m.name).upper() for m in s.execute(select(Machine)).scalars()}
    seen, m_new, m_exist, m_dup = set(), [], 0, []
    for m in parsed["machines"]:
        key = m["name"].upper()
        if key in seen:
            m_dup.append(m); continue
        seen.add(key)
        if key in exist_m:
            m_exist += 1
        else:
            m_new.append(m)

    return {
        "rows": len(rows) + len(parsed["invalid"]), "valid_rows": len(rows), "codes": len(groups),
        "new_count": len(new), "same_count": len(same), "changed_count": len(changed),
        "stock_changed_count": sum(1 for c in changed if c["stock"]),
        "dup_in_file_count": len(in_file), "dup_rows_extra": sum(len(d["rows"]) - 1 for d in in_file),
        "possible_dup_count": len(possible), "invalid_count": len(parsed["invalid"]),
        "machines_new": len(m_new), "machines_existing": m_exist, "machines_dup_in_file": len(m_dup),
        "new": new[:MAX_LIST], "same": same[:MAX_LIST], "changed": changed[:MAX_LIST],
        "dup_in_file": in_file[:MAX_LIST], "possible_dup": possible[:MAX_LIST],
        "invalid": parsed["invalid"][:MAX_LIST], "machines_new_list": [m["name"] for m in m_new][:MAX_LIST],
        "truncated": any(len(x) > MAX_LIST for x in (new, same, changed, in_file, possible, parsed["invalid"])),
    }


def wipe_inventory(s):
    """Children before parents (PostgreSQL enforces foreign keys)."""
    s.query(ImageBlob).delete(synchronize_session=False)
    for m in (ReturnLog, POLine, PurchaseOrder, ReqLine, Requisition, StockTxn, Item, Machine):
        s.query(m).delete(synchronize_session=False)
    s.flush()


def apply(s, parsed, mode="update", dup_rows="first", update_stock=False, skip_codes=(), user_id=None):
    if mode not in ("new_only", "update", "replace"):
        raise ImportError_("unknown mode")
    if dup_rows not in ("first", "sum"):
        raise ImportError_("unknown duplicate handling")
    skip = {code_key(c) for c in skip_codes}
    if mode == "replace":
        wipe_inventory(s)
    existing = {code_key(i.item_code): i for i in s.execute(select(Item)).scalars()}
    res = {"added": 0, "updated": 0, "stock_set": 0, "unchanged": 0, "skipped_existing": 0,
           "skipped_dup_rows": 0, "summed_codes": 0, "skipped_by_user": 0, "machines_added": 0}
    now = datetime.utcnow()
    for k, grp in _group(parsed["rows"]).items():
        if k in skip:
            res["skipped_by_user"] += 1; continue
        if len(grp) > 1:
            if dup_rows == "first":
                res["skipped_dup_rows"] += len(grp) - 1
            else:
                res["summed_codes"] += 1
        r = _merge_group(grp, dup_rows)
        it = existing.get(k)
        if it is None:
            it = Item(item_code=r["item_code"], updated_at=now, source_sheet="import",
                      **{f: r[f] for f in TEXT_FIELDS + NUM_FIELDS})
            it.max_level = max(r["reorder_point"], r["on_hand"])
            s.add(it); s.flush()
            note = "นำเข้าจาก Template" + (f" (รวม {len(grp)} แถว)" if len(grp) > 1 and dup_rows == "sum" else "")
            s.add(StockTxn(item_id=it.id, item_code=it.item_code, txn_type="receive", qty=it.on_hand,
                           balance_after=it.on_hand, ref="IMPORT", note=note, user_id=user_id))
            existing[k] = it
            res["added"] += 1
            continue
        if mode == "new_only":
            res["skipped_existing"] += 1; continue
        changed = False
        for f in DETAIL_FIELDS:
            if f in r["_has"] and not _same(getattr(it, f), r[f]):
                setattr(it, f, r[f]); changed = True
        if update_stock and "on_hand" in r["_has"]:
            cur = s.execute(select(Item.on_hand).where(Item.id == it.id).with_for_update()).scalar_one() or 0
            diff = round(r["on_hand"] - cur, 6)
            if diff:
                bal = change_stock(s, it.id, diff)
                # a counted stock figure becomes the new starting balance
                s.add(StockTxn(item_id=it.id, item_code=it.item_code, txn_type="adjust", qty=diff,
                               balance_after=bal, ref="IMPORT", note="ตั้งยอดคงเหลือตาม Template", user_id=user_id))
                res["stock_set"] += 1; changed = True
        if changed:
            it.updated_at = now; res["updated"] += 1
        else:
            res["unchanged"] += 1
    exist_m = {norm_text(m.name).upper() for m in s.execute(select(Machine)).scalars()}
    for m in parsed["machines"]:
        if m["name"].upper() not in exist_m:
            s.add(Machine(name=m["name"])); exist_m.add(m["name"].upper()); res["machines_added"] += 1
    s.commit()
    res["invalid_skipped"] = len(parsed["invalid"])
    return res
